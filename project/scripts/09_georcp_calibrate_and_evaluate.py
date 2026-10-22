# =============================================================================
# RVO-ME STEP 09
# GeoRCP-MIRAGE — VALIDATION CALIBRATION + ONE-TIME HELD-OUT EVALUATION
#
# RUN 1: validation-only threshold calibration, freeze evaluation protocol, exit.
# RUN 2: verify unchanged script/protocol, then evaluate all five held-out folds.
# =============================================================================
from __future__ import annotations

import ast, csv, gc, gzip, hashlib, importlib.util, json, math, re, sys, time
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import cv2
import numpy as np
import torch
import torch.nn.functional as F
from torch.utils.data import Dataset, DataLoader
from scipy.optimize import linear_sum_assignment
from openpyxl import load_workbook

SCRIPT_REVISION = "STEP09_GEOMETRY_EVAL_V1_20260822"
ROOT = Path(r"D:\RVO_ME")
DATASET = ROOT / "RVO-Lesion"
PROJECT = ROOT / "project"
DEMOGRAPHICS_XLSX = DATASET / "Demographics of the participants.xlsx"
STEP07_SCRIPT = PROJECT / "scripts" / "07_georcp_preflight_v2_ampfix_py38compat.py"
SPLIT_ROOT = PROJECT / "splits" / "patient_5fold_v1"
TARGET_ROOT = PROJECT / "targets" / "geometry_v3"
TARGET_DIR = TARGET_ROOT / "targets"
G3_COMPLETE = TARGET_ROOT / "GEOMETRY_V3_COMPLETE.json"
MIRAGE_WRAPPER = PROJECT / "foundation_models" / "mirage_base_preflight_v1" / "official_mirage_hf.py"
TRAIN_ROOT = PROJECT / "training" / "georcp_v1" / "definitive_v2"
TRAIN_MANIFEST = TRAIN_ROOT / "FIVEFOLD_TRAINING_MANIFEST.json"
GLOBAL_PROTOCOL_FILE = PROJECT / "design" / "geometry_aware_rcp_v1" / "RTX4080_ENGINEERING_FROZEN_v2.json"
EVAL_ROOT = PROJECT / "evaluation" / "georcp_v1" / "trackB_fivefold_v1"
CALIBRATION_ROOT = EVAL_ROOT / "calibration"
EVAL_PROTOCOL_FILE = EVAL_ROOT / "STEP09_EVALUATION_PROTOCOL_FROZEN.json"
HELDOUT_STARTED_FILE = EVAL_ROOT / "HELDOUT_EVALUATION_STARTED.json"
HELDOUT_COMPLETE_FILE = EVAL_ROOT / "STEP09_HELDOUT_RESULTS_COMPLETE.json"
PER_IMAGE_ALL_CSV = EVAL_ROOT / "heldout_per_image_metrics.csv"
SUMMARY_CSV = EVAL_ROOT / "heldout_summary.csv"
PROJECT_LOG = PROJECT / "PROJECT_LOG.md"

EXPECTED_DESIGN_SHA = "9f1586d102523009443e2b88dca54a9ca8fa5c983f08fdde837184d436a95801"
EXPECTED_G3_SHA = "521859acb4c5389eb4244298f755420f1339b7e4c4bedcfac7c64ddcfe98ce59"
EXPECTED_IMPLEMENTATION_SPEC_SHA = "b5e59901eb67190d5924698ca53165d30d1bccf45028a7454d9ca98147149422"
EXPECTED_SOURCE_BUNDLE_SHA = "04e7da4fe311b756ad375a6c241530d0572a3cf1b9c71a70af2e7e7e24524412"
EXPECTED_MIRAGE_WRAPPER_SHA = "07fa63e2766b67dbc0c07f0d9d29db1562f707425ef2fbf90a688f62a94e6ede"
EXPECTED_TRACKB_SPLIT_SHA = "7e7c2ddf51054bf998b2d750b50c74983b9bc5b751b762213bc96343ffe01346"
EXPECTED_GLOBAL_PROTOCOL_SHA = "8640e7b4039cfc5f6fcd45269aaf9003a337a090f3af5a8aa787c9dcd9a0491e"
EXPECTED_IMAGES = 3012
EXPECTED_VAL_IMAGES = {1:303, 2:270, 3:308, 4:342, 5:362}
EXPECTED_TEST_IMAGES = {1:599, 2:606, 3:598, 4:631, 5:578}
EXPECTED_TEST_PATIENTS = 26
REGION_NAMES = {1:"SRF", 2:"IRF"}
CURVE_NAMES = {0:"ELM", 1:"EZ"}

NH, NW = 380, 570
PT, PB, PL, PR = 2, 2, 3, 3
H, W = 384, 576
HF_STRIDE = 2
EVAL_BATCH = 7
BOUNDARY_TOLERANCE_PX = 2.0
CURVE_THRESHOLD_CANDIDATES = [round(float(x),2) for x in np.arange(0.10,0.901,0.05)]
HF_THRESHOLD_CANDIDATES = [round(float(x),2) for x in np.arange(0.05,0.951,0.05)]
HF_LOCAL_MAX_KERNEL = 3
HF_TOPK_PER_IMAGE = 100
HF_MIN_PEAK_SCORE = 0.01
HF_RADII_PX = (5.0, 10.0)
FROC_RADIUS_PX = 10.0
FROC_FP_PER_IMAGE = (0.25,0.5,1.0,2.0,4.0,8.0)
HF_BOX_SIZE_PX = 10.0
HF_AP_IOU = 0.50
BOOTSTRAP_REPLICATES = 2000
BOOTSTRAP_SEED = 20260822
EPS = 1e-12

SOURCE_BUNDLE_NAMES = [
    "build_mirage_class","load_mirage","LN2d","CNBlock","LocalEncoder",
    "GlobalFusion","TaskInteraction","Reassemble","TaskFusion","FullTrunk",
    "RegionHead","CurveHead","PointHead","GeoRCP","masked_mean","binary_dice",
    "foreground_dice","centernet_focal","GeoRCPLoss","mirage_norm","load_sample",
    "optimizer_for","build_model",
]


def load_json(path: Path):
    if not path.exists(): raise FileNotFoundError(path)
    return json.loads(path.read_text(encoding="utf-8"))


def save_json(path: Path, obj: Any):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, indent=2, ensure_ascii=False, sort_keys=True), encoding="utf-8")


def save_json_gz(path: Path, obj: Any):
    path.parent.mkdir(parents=True, exist_ok=True)
    with gzip.open(path, "wt", encoding="utf-8") as f: json.dump(obj, f, ensure_ascii=False, separators=(",",":"))


def load_json_gz(path: Path):
    with gzip.open(path, "rt", encoding="utf-8") as f: return json.load(f)


def sha_file(path: Path):
    h=hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda:f.read(1024*1024), b""): h.update(chunk)
    return h.hexdigest()


def stable_sha(obj: Any):
    return hashlib.sha256(json.dumps(obj,sort_keys=True,ensure_ascii=False,separators=(",",":")).encode()).hexdigest()


def current_script_sha(): return sha_file(Path(__file__).resolve())

def hash_stems(stems): return hashlib.sha256("\n".join(sorted(stems)).encode()).hexdigest()

def reset_cuda():
    gc.collect()
    if torch.cuda.is_available(): torch.cuda.empty_cache(); torch.cuda.synchronize()

def safe_div(a,b): return float("nan") if abs(float(b))<=EPS else float(a)/float(b)

def finite_values(values): return [float(v) for v in values if np.isfinite(v)]

def f1_counts(tp,fp,fn):
    p=safe_div(tp,tp+fp); r=safe_div(tp,tp+fn)
    f=0.0 if (not np.isfinite(p) or not np.isfinite(r) or p+r<=EPS) else 2*p*r/(p+r)
    return {"precision":float(p),"recall":float(r),"f1":float(f)}


def append_log_once(marker,text):
    if not PROJECT_LOG.exists(): return
    old=PROJECT_LOG.read_text(encoding="utf-8",errors="ignore")
    if marker in old: return
    with PROJECT_LOG.open("a",encoding="utf-8") as f: f.write("\n\n{}\n{}\n".format(marker,text))


def module_from_file(name,path):
    spec=importlib.util.spec_from_file_location(name,path)
    if spec is None or spec.loader is None: raise RuntimeError("Could not import {}".format(path))
    mod=importlib.util.module_from_spec(spec); sys.modules[name]=mod; spec.loader.exec_module(mod); return mod


def source_bundle_sha(path):
    text=path.read_text(encoding="utf-8").replace("\r\n","\n")
    tree=ast.parse(text,filename=str(path)); lines=text.splitlines(); wanted=set(SOURCE_BUNDLE_NAMES); chunks=[]; found=[]
    for node in tree.body:
        if isinstance(node,(ast.FunctionDef,ast.AsyncFunctionDef,ast.ClassDef)) and node.name in wanted:
            if not hasattr(node,"end_lineno"): raise RuntimeError("AST lacks end_lineno")
            chunks.append("## {}\n{}\n".format(node.name,"\n".join(lines[node.lineno-1:node.end_lineno]))); found.append(node.name)
    missing=sorted(wanted-set(found))
    if missing: raise RuntimeError("Missing source definitions: "+", ".join(missing))
    return hashlib.sha256("\n".join(chunks).encode()).hexdigest()


def load_exact_step07():
    s=source_bundle_sha(STEP07_SCRIPT)
    if s!=EXPECTED_SOURCE_BUNDLE_SHA: raise RuntimeError("GeoRCP source-bundle SHA mismatch")
    mod=module_from_file("georcp_step07_eval",STEP07_SCRIPT)
    if mod.implementation_spec().get("implementation_spec_sha256")!=EXPECTED_IMPLEMENTATION_SPEC_SHA: raise RuntimeError("Implementation SHA mismatch")
    return mod,s


def verify_training_manifest(source_sha):
    if not torch.cuda.is_available(): raise RuntimeError("CUDA required")
    if "RTX 4080" not in torch.cuda.get_device_name(0): raise RuntimeError("RTX 4080 required")
    if not sys.version.startswith("3.8"): raise RuntimeError("Validated Python 3.8.x required")
    if torch.__version__!="2.4.1+cu121" or str(torch.version.cuda)!="12.1": raise RuntimeError("Validated PyTorch/CUDA environment mismatch")
    if sha_file(MIRAGE_WRAPPER)!=EXPECTED_MIRAGE_WRAPPER_SHA: raise RuntimeError("MIRAGE wrapper SHA mismatch")
    if source_sha!=EXPECTED_SOURCE_BUNDLE_SHA: raise RuntimeError("Source SHA mismatch")
    if load_json(G3_COMPLETE).get("geometry_v3_master_content_sha256")!=EXPECTED_G3_SHA: raise RuntimeError("geometry_v3 SHA mismatch")
    if len(list(TARGET_DIR.glob("*.npz")))!=EXPECTED_IMAGES: raise RuntimeError("geometry_v3 target count mismatch")
    if load_json(GLOBAL_PROTOCOL_FILE).get("protocol_sha256")!=EXPECTED_GLOBAL_PROTOCOL_SHA: raise RuntimeError("Global protocol SHA mismatch")
    m=load_json(TRAIN_MANIFEST)
    if m.get("status")!="ALL_FIVE_FOLDS_TRAINING_COMPLETE" or m.get("global_protocol_sha256")!=EXPECTED_GLOBAL_PROTOCOL_SHA: raise RuntimeError("Five-fold manifest invalid")
    if m.get("selected_protocol")!="NO_AUGMENTATION" or m.get("outer_test_resolved") or m.get("outer_test_evaluated"): raise RuntimeError("Training manifest test/protocol policy mismatch")
    for fold in range(1,6):
        e=m["folds"][str(fold)]
        if e.get("status")!="COMPLETE": raise RuntimeError("Fold {} incomplete".format(fold))
        cp=Path(e["best_checkpoint"]); rp=Path(e["completion_report"])
        if not cp.exists() or not rp.exists(): raise RuntimeError("Fold {} checkpoint/report missing".format(fold))
        r=load_json(rp)
        if r.get("status")!="COMPLETE" or r.get("active_target_master_sha256")!=EXPECTED_G3_SHA or r.get("implementation_spec_sha256")!=EXPECTED_IMPLEMENTATION_SPEC_SHA or r.get("source_bundle_sha256")!=EXPECTED_SOURCE_BUNDLE_SHA or r.get("outer_test_evaluated"): raise RuntimeError("Fold {} completion integrity mismatch".format(fold))
        psha=r.get("protocol_sha256") if fold==1 else r.get("global_protocol_sha256")
        if psha!=EXPECTED_GLOBAL_PROTOCOL_SHA: raise RuntimeError("Fold {} protocol mismatch".format(fold))
    return m


def canonical_id(v):
    if v is None: return ""
    if isinstance(v,float) and not math.isnan(v) and v.is_integer(): return str(int(v))
    s=str(v).strip()
    return s[:-2] if re.fullmatch(r"\d+\.0",s) else s


def load_eye_to_patient():
    wb=load_workbook(DEMOGRAPHICS_XLSX,read_only=True,data_only=True); ws=wb[wb.sheetnames[0]]; rows=ws.iter_rows(values_only=True)
    hdr=next(rows); norm={re.sub(r"[^a-z0-9]+","",str(v).lower()):i for i,v in enumerate(hdr) if v is not None}
    pc=next((norm[k] for k in ("patientid","patient") if k in norm),None); ec=next((norm[k] for k in ("eyeid","eye") if k in norm),None)
    if pc is None or ec is None: raise RuntimeError("Patient ID / Eye ID columns not found")
    out={}
    for row in rows:
        if max(pc,ec)>=len(row): continue
        p,e=canonical_id(row[pc]),canonical_id(row[ec])
        if p and e:
            if e in out and out[e]!=p: raise RuntimeError("Eye maps to multiple patients")
            out[e]=p
    if len(set(out.values()))!=130: raise RuntimeError("Expected 130 patients; found {}".format(len(set(out.values()))))
    return out


def patient_for_stem(stem,mapping):
    eye=canonical_id(stem.split("_",1)[0]); p=mapping.get(eye)
    if p is None: raise RuntimeError("No patient mapping for {}".format(stem))
    return p

STEM_COLUMNS={"stem","image","image_id","filename","file","name"}; ROLE_COLUMNS={"role","split","subset","set"}; FOLD_COLUMNS={"fold","fold_id","outer_fold"}
def norm_fold(v):
    m=re.search(r"(\d+)",str(v).strip().lower()); return int(m.group(1)) if m else None


def resolve_split_role(fold,role,available_stems):
    if role not in {"val","test"}: raise RuntimeError("Step09 resolves VAL/TEST only")
    available=set(available_stems); found=set(); files=sorted(p for p in SPLIT_ROOT.rglob("*") if p.is_file() and p.suffix.lower() in {".txt",".csv",".json"}); toks={"fold_{}".format(fold),"fold{}".format(fold),"/{}/".format(fold)}
    for path in files:
        low=str(path).lower().replace("\\","/")
        if path.suffix.lower()==".txt" and any(t in low for t in toks) and role in path.name.lower():
            for token in re.split(r"[\s,;]+",path.read_text(encoding="utf-8",errors="ignore")):
                s=Path(token.strip("\"'")).stem
                if s in available: found.add(s)
    for path in files:
        if path.suffix.lower()!=".csv": continue
        low=str(path).lower().replace("\\","/"); pf=fold if any(t in low for t in toks) else None
        try:
            with path.open("r",encoding="utf-8-sig",newline="") as f:
                rd=csv.DictReader(f); cmap={c.lower():c for c in (rd.fieldnames or [])}; sc=next((cmap[c] for c in STEM_COLUMNS if c in cmap),None); rc=next((cmap[c] for c in ROLE_COLUMNS if c in cmap),None); fc=next((cmap[c] for c in FOLD_COLUMNS if c in cmap),None)
                if sc is None: continue
                for row in rd:
                    rr=str(row.get(rc,"")).strip().lower() if rc else (role if role in path.name.lower() else "")
                    rf=norm_fold(row.get(fc,"")) if fc else pf
                    if rr==role and rf==fold:
                        s=Path(str(row.get(sc,""))).stem
                        if s in available: found.add(s)
        except Exception: pass
    def walk(obj,inf=None,inr=None):
        if isinstance(obj,dict):
            low={str(k).lower():v for k,v in obj.items()}; lf,lr=inf,inr
            for k in FOLD_COLUMNS:
                if k in low: lf=norm_fold(low[k])
            for k in ROLE_COLUMNS:
                if k in low: lr=str(low[k]).strip().lower()
            sv=next((low[k] for k in STEM_COLUMNS if k in low),None)
            if lf==fold and lr==role and sv is not None:
                s=Path(str(sv)).stem
                if s in available: found.add(s)
            for v in obj.values(): walk(v,lf,lr)
        elif isinstance(obj,list):
            for v in obj: walk(v,inf,inr)
    for path in files:
        if path.suffix.lower()==".json":
            try: walk(load_json(path))
            except Exception: pass
    stems=sorted(found); expected=(EXPECTED_VAL_IMAGES if role=="val" else EXPECTED_TEST_IMAGES)[fold]
    if len(stems)!=expected: raise RuntimeError("Fold {} {} expected {}, found {}".format(fold,role,expected,len(stems)))
    return stems


class GeoRCPDataset(Dataset):
    def __init__(self,stems,images,step07): self.stems=list(stems); self.images=images; self.step07=step07
    def __len__(self): return len(self.stems)
    def __getitem__(self,i):
        s=self.stems[i]; x=self.step07.load_sample(s,self.images[s],TARGET_DIR); x["_stem"]=s; return x


def collate(items):
    stems=[x["_stem"] for x in items]; keys=[k for k in items[0] if k!="_stem"]; b={k:torch.stack([x[k] for x in items]) for k in keys}; b["_stems"]=stems; return b

def loader(ds): return DataLoader(ds,batch_size=EVAL_BATCH,shuffle=False,num_workers=0,pin_memory=False,drop_last=False,collate_fn=collate)


def torch_load(path,device):
    try: return torch.load(path,map_location=device,weights_only=False)
    except TypeError: return torch.load(path,map_location=device)


def build_fold_model(fold,checkpoint,step07,device):
    reset_cuda(); model,loss=step07.build_model(device); st=torch_load(checkpoint,device)
    if st.get("active_target_master_sha256")!=EXPECTED_G3_SHA or st.get("implementation_spec_sha256")!=EXPECTED_IMPLEMENTATION_SPEC_SHA or st.get("source_bundle_sha256")!=EXPECTED_SOURCE_BUNDLE_SHA: raise RuntimeError("Fold {} checkpoint integrity mismatch".format(fold))
    psha=st.get("protocol_sha256") if fold==1 else st.get("global_protocol_sha256")
    if psha!=EXPECTED_GLOBAL_PROTOCOL_SHA: raise RuntimeError("Fold {} checkpoint protocol mismatch".format(fold))
    model.load_state_dict(st["model"],strict=True); model.eval(); del loss; return model


def gt_hf(sample):
    m=sample["hf_offset_mask"].numpy(); off=sample["hf_offset"].numpy(); ys,xs=np.where(m>0); pts=[]
    for cy,cx in zip(ys,xs):
        x=(cx+float(off[0,cy,cx]))*HF_STRIDE-PL; y=(cy+float(off[1,cy,cx]))*HF_STRIDE-PT
        if -1e-4<=x<NW+1e-4 and -1e-4<=y<NH+1e-4: pts.append([x,y])
    return np.asarray(pts,dtype=np.float32).reshape(-1,2)


def decode_hf(logit,offset):
    p=torch.sigmoid(logit); mx=F.max_pool2d(p,HF_LOCAL_MAX_KERNEL,1,HF_LOCAL_MAX_KERNEL//2); keep=(p>=mx-1e-12)&(p>=HF_MIN_PEAK_SCORE); scores=p[keep]; idx=keep.nonzero(as_tuple=False)
    if scores.numel()==0: return []
    order=torch.argsort(scores,descending=True)[:HF_TOPK_PER_IMAGE]; out=[]
    for j in order.tolist():
        _,cy,cx=idx[j].tolist(); score=float(scores[j].cpu()); dx=float(offset[0,cy,cx].cpu()); dy=float(offset[1,cy,cx].cpu()); x=(cx+dx)*HF_STRIDE-PL; y=(cy+dy)*HF_STRIDE-PT
        if 0<=x<NW and 0<=y<NH: out.append({"score":score,"x":float(x),"y":float(y)})
    return out


def gt_curve_y(center):
    c=center[:,PT:PT+NH,PL:PL+NW]; sums=c.sum(axis=1); yy=np.arange(NH,dtype=np.float32)[None,:,None]; num=(c*yy).sum(axis=1); out=np.zeros((2,NW),np.float32); v=sums>1e-8; out[v]=num[v]/sums[v]; return out


def pred_curve_y(logit):
    c=logit[:,PT:PT+NH,PL:PL+NW]; p=torch.softmax(c,dim=1); yy=torch.arange(NH,device=p.device,dtype=p.dtype).view(1,NH,1); return (p*yy).sum(1).float().cpu().numpy()


def match_points(pred,gt,radius):
    npred,ngt=len(pred),len(gt)
    if npred==0 and ngt==0: return 0,0,0,np.zeros(0,np.float32)
    if npred==0: return 0,0,ngt,np.zeros(0,np.float32)
    if ngt==0: return 0,npred,0,np.zeros(0,np.float32)
    d=np.sqrt(((pred[:,None,:]-gt[None,:,:])**2).sum(2)); cost=np.where(d<=radius,d,1e6); rr,cc=linear_sum_assignment(cost); ds=np.asarray([d[r,c] for r,c in zip(rr,cc) if d[r,c]<=radius],np.float32); tp=len(ds); return tp,npred-tp,ngt-tp,ds


def binary_boundary(mask):
    m=mask.astype(np.uint8)
    if not m.any(): return np.zeros_like(m,bool)
    e=cv2.erode(m,np.ones((3,3),np.uint8),iterations=1); return (m>0)&(e==0)


def boundary_stats(pred,gt):
    pb,gb=binary_boundary(pred),binary_boundary(gt); r=int(math.ceil(BOUNDARY_TOLERANCE_PX)); ker=cv2.getStructuringElement(cv2.MORPH_ELLIPSE,(2*r+1,2*r+1)); pc,gc=int(pb.sum()),int(gb.sum())
    ph=int((pb & (cv2.dilate(gb.astype(np.uint8),ker)>0)).sum()) if pc else 0; gh=int((gb & (cv2.dilate(pb.astype(np.uint8),ker)>0)).sum()) if gc else 0
    if pc==0 and gc==0: hd=float("nan")
    elif pc==0 or gc==0: hd=math.sqrt((NW-1)**2+(NH-1)**2)
    else:
        dg=cv2.distanceTransform((~gb).astype(np.uint8),cv2.DIST_L2,5); dp=cv2.distanceTransform((~pb).astype(np.uint8),cv2.DIST_L2,5); hd=float(np.percentile(np.concatenate([dg[pb],dp[gb]]),95))
    return ph,pc,gh,gc,hd


def segments(x):
    a=x.astype(np.uint8); return int(((a==1)&np.concatenate([[True],a[:-1]==0])).sum()) if len(a) else 0


def infer_records(fold,stems,checkpoint,step07,images,device,regions=False):
    model=build_fold_model(fold,checkpoint,step07,device); ds=GeoRCPDataset(stems,images,step07); ld=loader(ds); recs=[]
    with torch.no_grad():
        for b in ld:
            im=b["image"].to(device)
            with torch.amp.autocast("cuda",dtype=torch.bfloat16): o=model(im)
            rp=o["region_logits"].argmax(1).cpu().numpy(); pp=torch.sigmoid(o["curve_presence_logits"])[:,:,PL:PL+NW].float().cpu().numpy()
            for i,stem in enumerate(b["_stems"]):
                s={k:b[k][i].cpu() for k in b if k!="_stems"}; r={"stem":stem,"curve_prob":pp[i],"curve_gt":s["curve_presence"].numpy()[:,PL:PL+NW]>0.5,"curve_valid":s["curve_column_valid"].numpy()[PL:PL+NW]>0.5,"curve_py":pred_curve_y(o["curve_centerline_logits"][i]),"curve_gy":gt_curve_y(s["curve_centerline"].numpy()),"hf_pred":decode_hf(o["hf_heatmap_logits"][i],o["hf_offset"][i]),"hf_gt":gt_hf(s)}
                if regions: r.update({"region_pred":rp[i,PT:PT+NH,PL:PL+NW],"region_gt":s["region_class"].numpy()[PT:PT+NH,PL:PL+NW],"region_valid":s["region_valid"].numpy()[PT:PT+NH,PL:PL+NW]>0.5})
                recs.append(r)
    del ld,ds,model; reset_cuda(); return recs


def calibrate_curve(recs):
    table=[]
    for t in CURVE_THRESHOLD_CANDIDATES:
        fs=[]; row={"threshold":t}
        for ch,name in CURVE_NAMES.items():
            tp=fp=fn=0
            for r in recs:
                v=r["curve_valid"]; g=r["curve_gt"][ch]&v; p=(r["curve_prob"][ch]>=t)&v; tp+=int((p&g).sum()); fp+=int((p&~g).sum()); fn+=int((~p&g).sum())
            f=f1_counts(tp,fp,fn)["f1"]; row[name+"_f1"]=f; fs.append(f)
        row["macro_f1"]=float(np.mean(fs)); table.append(row)
    best=max(table,key=lambda r:(r["macro_f1"],-abs(r["threshold"]-0.5)))
    return {"selected_threshold":best["threshold"],"selected_score":best["macro_f1"],"selection_metric":"macro ELM/EZ column-presence F1","tie_break":"closest to 0.5","candidate_table":table}


def hf_xy_at(r,t):
    a=[[d["x"],d["y"]] for d in r["hf_pred"] if d["score"]>=t]; return np.asarray(a,np.float32).reshape(-1,2)


def calibrate_hf(recs):
    table=[]
    for t in HF_THRESHOLD_CANDIDATES:
        fs={}
        for rad in HF_RADII_PX:
            tp=fp=fn=0
            for r in recs:
                a,b,c,_=match_points(hf_xy_at(r,t),r["hf_gt"],rad); tp+=a; fp+=b; fn+=c
            fs[rad]=f1_counts(tp,fp,fn)["f1"]
        score=float(np.mean([fs[5.0],fs[10.0]])); table.append({"threshold":t,"f1_5px":fs[5.0],"f1_10px":fs[10.0],"selection_score":score})
    best=max(table,key=lambda r:(r["selection_score"],r["threshold"]))
    return {"selected_threshold":best["threshold"],"selected_score":best["selection_score"],"selection_metric":"mean pooled HF F1@5px and F1@10px","tie_break":"higher threshold","candidate_table":table}


def run_calibration(manifest,step07,images,available,device):
    CALIBRATION_ROOT.mkdir(parents=True,exist_ok=True); out={}
    for fold in range(1,6):
        print("\nCALIBRATION FOLD {}".format(fold)); stems=resolve_split_role(fold,"val",available); cp=Path(manifest["folds"][str(fold)]["best_checkpoint"]); recs=infer_records(fold,stems,cp,step07,images,device,regions=False); c=calibrate_curve(recs); h=calibrate_hf(recs); res={"fold":fold,"validation_images":len(stems),"validation_stem_sha256":hash_stems(stems),"curve":c,"hf":h}; save_json(CALIBRATION_ROOT/("fold_{}_calibration.json".format(fold)),res); out[str(fold)]=res; print("  curve threshold {:.2f}, score {:.5f}".format(c["selected_threshold"],c["selected_score"])); print("  HF threshold {:.2f}, score {:.5f}".format(h["selected_threshold"],h["selected_score"])); del recs
    return out


def make_protocol(cal):
    p={"status":"FROZEN_BEFORE_HELDOUT_TEST","script_revision":SCRIPT_REVISION,"script_sha256":current_script_sha(),"scientific_design_sha256":EXPECTED_DESIGN_SHA,"active_target_master_sha256":EXPECTED_G3_SHA,"implementation_spec_sha256":EXPECTED_IMPLEMENTATION_SPEC_SHA,"source_bundle_sha256":EXPECTED_SOURCE_BUNDLE_SHA,"mirage_wrapper_sha256":EXPECTED_MIRAGE_WRAPPER_SHA,"trackb_split_master_sha256":EXPECTED_TRACKB_SPLIT_SHA,"global_training_protocol_sha256":EXPECTED_GLOBAL_PROTOCOL_SHA,"calibration_scope":"fold-specific validation only","heldout_test_used_for_calibration":False,"region":{"decode":"argmax bg/SRF/IRF, native crop","overlap":"exclude region_valid=0","boundary_source":"derived from semantic mask","boundary_f1_tolerance_px":BOUNDARY_TOLERANCE_PX,"hd95":"symmetric native-pixel boundary HD95"},"curve":{"y_decode":"vertical softmax expected-y over native rows","presence":"fold-specific validation threshold","no_gap_fill":True,"no_connectivity_postprocessing":True,"localization":"absolute y error where GT and prediction both present","gap_bridge":"predicted presence on internal GT-absent columns","secondary":"1-pixel centerline Dice"},"hf":{"activation":"sigmoid","local_max_kernel":HF_LOCAL_MAX_KERNEL,"topk":HF_TOPK_PER_IMAGE,"min_candidate_score":HF_MIN_PEAK_SCORE,"matching":"Hungarian within radius","radii_px":list(HF_RADII_PX),"froc_radius_px":FROC_RADIUS_PX,"froc_fp_per_image":list(FROC_FP_PER_IMAGE),"ap_box_size_px":HF_BOX_SIZE_PX,"ap_iou":HF_AP_IOU},"threshold_calibration":cal,"patient_bootstrap":{"replicates":BOOTSTRAP_REPLICATES,"seed":BOOTSTRAP_SEED,"ci":"95% percentile","unit":"patient"},"test_policy":"one frozen pass per fold; no test-derived tuning"}; p["evaluation_protocol_sha256"]=stable_sha(p); return p


def eval_region(predc,gtc,valid,cid):
    p=(predc==cid)&valid; g=(gtc==cid)&valid; inter=int((p&g).sum()); pn=int(p.sum()); gn=int(g.sum()); union=int((p|g).sum()); ph,pc,gh,gc,hd=boundary_stats(p,g); return {"inter":inter,"pred":pn,"gt":gn,"union":union,"boundary_pred_hit":ph,"boundary_pred_count":pc,"boundary_gt_hit":gh,"boundary_gt_count":gc,"hd95":hd}


def symmetric_curve_hd(pred,gt):
    if len(pred)==0 and len(gt)==0: return float("nan")
    if len(pred)==0 or len(gt)==0: return math.sqrt((NW-1)**2+(NH-1)**2)
    d=np.sqrt(((pred[:,None,:]-gt[None,:,:])**2).sum(2)); return float(np.percentile(np.concatenate([d.min(1),d.min(0)]),95))


def eval_curve(prob,py,gp,gy,valid,t):
    g=gp&valid; p=(prob>=t)&valid; tp=int((p&g).sum()); fp=int((p&~g).sum()); fn=int((~p&g).sum()); both=p&g; err=np.abs(py[both]-gy[both]).astype(np.float32); xp=np.where(p)[0].astype(np.float32); xg=np.where(g)[0].astype(np.float32); pp=np.stack([xp,py[p].astype(np.float32)],1) if len(xp) else np.zeros((0,2),np.float32); gg=np.stack([xg,gy[g].astype(np.float32)],1) if len(xg) else np.zeros((0,2),np.float32); hd=symmetric_curve_hd(pp,gg)
    if g.any():
        inds=np.where(g)[0]; internal=np.zeros_like(g,bool); internal[inds[0]:inds[-1]+1]=True; gap=internal&~g&valid; bt=int(gap.sum()); bp=int((p&gap).sum())
    else: bt=bp=0
    exact=int((both&(np.rint(py).astype(int)==np.rint(gy).astype(int))).sum())
    return {"tp":tp,"fp":fp,"fn":fn,"loc_sum":float(err.sum()),"loc_count":int(len(err)),"loc_median":float(np.median(err)) if len(err) else float("nan"),"hd95":hd,"bridge_pred":bp,"bridge_total":bt,"pred_points":int(p.sum()),"gt_points":int(g.sum()),"centerline_exact_inter":exact,"gt_fragments":segments(g),"pred_fragments":segments(p)},err


def eval_hf(dets,gt,t):
    pred=np.asarray([[d["x"],d["y"]] for d in dets if d["score"]>=t],np.float32).reshape(-1,2); out={"gt_count":len(gt),"pred_count":len(pred),"count_abs_error":abs(len(pred)-len(gt))}; errs={}
    for rad in HF_RADII_PX:
        tp,fp,fn,ds=match_points(pred,gt,rad); tag=str(int(rad)); out.update({"tp_"+tag:tp,"fp_"+tag:fp,"fn_"+tag:fn,"loc_sum_"+tag:float(ds.sum()),"loc_count_"+tag:len(ds),"loc_median_"+tag:float(np.median(ds)) if len(ds) else float("nan")}); errs[tag]=ds
    return out,errs


def box_iou(p,g):
    h=HF_BOX_SIZE_PX/2; px1,py1,px2,py2=p[0]-h,p[1]-h,p[0]+h,p[1]+h; gx1,gy1,gx2,gy2=g[0]-h,g[1]-h,g[0]+h,g[1]+h; iw=max(0,min(px2,gx2)-max(px1,gx1)); ih=max(0,min(py2,gy2)-max(py1,gy1)); inter=iw*ih; return safe_div(inter,2*HF_BOX_SIZE_PX*HF_BOX_SIZE_PX-inter)


def ap50(payload):
    gt={x["stem"]:{"pts":[tuple(p) for p in x["gt"]],"used":[False]*len(x["gt"])} for x in payload}; ng=sum(len(x["gt"]) for x in payload); preds=sorted([(float(d["score"]),x["stem"],(float(d["x"]),float(d["y"]))) for x in payload for d in x["pred"]],reverse=True,key=lambda z:z[0]); tp=[]; fp=[]
    if ng==0: return float("nan")
    for _,stem,p in preds:
        info=gt[stem]; best=-1; bi=None
        for i,g in enumerate(info["pts"]):
            if info["used"][i]: continue
            z=box_iou(p,g)
            if z>best: best=z; bi=i
        if bi is not None and best>=HF_AP_IOU: info["used"][bi]=True; tp.append(1.0); fp.append(0.0)
        else: tp.append(0.0); fp.append(1.0)
    if not tp: return 0.0
    t=np.cumsum(tp); f=np.cumsum(fp); r=t/ng; p=t/np.maximum(t+f,EPS); mr=np.concatenate([[0],r,[1]]); mp=np.concatenate([[0],p,[0]])
    for i in range(len(mp)-1,0,-1): mp[i-1]=max(mp[i-1],mp[i])
    idx=np.where(mr[1:]!=mr[:-1])[0]; return float(np.sum((mr[idx+1]-mr[idx])*mp[idx+1]))


def froc(payload):
    curve=[]; ng=sum(len(x["gt"]) for x in payload); ni=len(payload)
    for t in np.linspace(0.01,0.99,99):
        tp=fp=fn=0
        for x in payload:
            p=np.asarray([[d["x"],d["y"]] for d in x["pred"] if d["score"]>=t],np.float32).reshape(-1,2); g=np.asarray(x["gt"],np.float32).reshape(-1,2); a,b,c,_=match_points(p,g,FROC_RADIUS_PX); tp+=a; fp+=b; fn+=c
        curve.append({"threshold":round(float(t),3),"sensitivity":safe_div(tp,ng),"fp_per_image":safe_div(fp,ni)})
    ops={str(v):(max([r["sensitivity"] for r in curve if np.isfinite(r["fp_per_image"]) and r["fp_per_image"]<=v]) if any(np.isfinite(r["fp_per_image"]) and r["fp_per_image"]<=v for r in curve) else 0.0) for v in FROC_FP_PER_IMAGE}
    return {"radius_px":FROC_RADIUS_PX,"operating_points":ops,"curve":curve}


def evaluate_test_fold(fold,stems,cp,th,step07,images,eye2pat,device,protocol_sha):
    d=EVAL_ROOT/"heldout"/("fold_{}".format(fold)); d.mkdir(parents=True,exist_ok=True); complete=d/"TEST_EVALUATION_COMPLETE.json"; csvp=d/"per_image_metrics.csv"; hfp=d/"hf_payload.json.gz"; npzp=d/"matched_distance_arrays.npz"
    if complete.exists():
        old=load_json(complete)
        if old.get("evaluation_protocol_sha256")!=protocol_sha: raise RuntimeError("Existing Fold {} result protocol mismatch".format(fold))
        print("  Fold {} already COMPLETE — skipping".format(fold)); return old
    recs=infer_records(fold,stems,cp,step07,images,device,regions=True); rows=[]; cerr={"ELM":[],"EZ":[]}; herr={"5":[],"10":[]}; payload=[]; ct=float(th["curve"]["selected_threshold"]); ht=float(th["hf"]["selected_threshold"])
    for r in recs:
        stem=r["stem"]; patient=patient_for_stem(stem,eye2pat); row={"fold":fold,"stem":stem,"patient_id":patient}
        for cid,name in REGION_NAMES.items():
            m=eval_region(r["region_pred"],r["region_gt"],r["region_valid"],cid)
            for k,v in m.items(): row["region_{}_{}".format(name.lower(),k)]=v
        for ch,name in CURVE_NAMES.items():
            m,e=eval_curve(r["curve_prob"][ch],r["curve_py"][ch],r["curve_gt"][ch],r["curve_gy"][ch],r["curve_valid"],ct)
            for k,v in m.items(): row["curve_{}_{}".format(name.lower(),k)]=v
            if len(e): cerr[name].append(e)
        m,es=eval_hf(r["hf_pred"],r["hf_gt"],ht)
        for k,v in m.items(): row["hf_"+k]=v
        for k,e in es.items():
            if len(e): herr[k].append(e)
        payload.append({"stem":stem,"patient_id":patient,"pred":r["hf_pred"],"gt":r["hf_gt"].tolist()}); rows.append(row)
    fields=sorted(rows[0]);
    with csvp.open("w",encoding="utf-8",newline="") as f:
        w=csv.DictWriter(f,fieldnames=fields); w.writeheader(); w.writerows(rows)
    save_json_gz(hfp,payload); arrays={"curve_elm_errors":np.concatenate(cerr["ELM"]) if cerr["ELM"] else np.zeros(0,np.float32),"curve_ez_errors":np.concatenate(cerr["EZ"]) if cerr["EZ"] else np.zeros(0,np.float32),"hf_5_errors":np.concatenate(herr["5"]) if herr["5"] else np.zeros(0,np.float32),"hf_10_errors":np.concatenate(herr["10"]) if herr["10"] else np.zeros(0,np.float32)}; np.savez_compressed(npzp,**arrays)
    pats={x["patient_id"] for x in rows}
    if len(pats)!=EXPECTED_TEST_PATIENTS: raise RuntimeError("Fold {} expected 26 patients, found {}".format(fold,len(pats)))
    out={"status":"COMPLETE","fold":fold,"test_images":len(stems),"test_patients":len(pats),"test_stem_sha256":hash_stems(stems),"evaluation_protocol_sha256":protocol_sha,"curve_presence_threshold":ct,"hf_score_threshold":ht,"per_image_csv":str(csvp),"hf_payload":str(hfp),"distance_arrays":str(npzp),"outer_test_evaluated":True,"post_test_tuning_allowed":False}; save_json(complete,out); return out


def read_rows(path):
    with path.open("r",encoding="utf-8-sig",newline="") as f: return list(csv.DictReader(f))
def fv(row,key,default=0.0):
    v=row.get(key)
    if v in (None,"","nan","NaN"): return float("nan")
    try: return float(v)
    except Exception: return float(default)


def agg_region(rows,name):
    p="region_{}_".format(name.lower()); inter=sum(fv(r,p+"inter") for r in rows); pred=sum(fv(r,p+"pred") for r in rows); gt=sum(fv(r,p+"gt") for r in rows); union=sum(fv(r,p+"union") for r in rows); ph=sum(fv(r,p+"boundary_pred_hit") for r in rows); pc=sum(fv(r,p+"boundary_pred_count") for r in rows); gh=sum(fv(r,p+"boundary_gt_hit") for r in rows); gc=sum(fv(r,p+"boundary_gt_count") for r in rows); bp=safe_div(ph,pc); br=safe_div(gh,gc); bf=0.0 if not np.isfinite(bp) or not np.isfinite(br) or bp+br<=EPS else 2*bp*br/(bp+br); hd=finite_values(fv(r,p+"hd95") for r in rows); return {"dice":safe_div(2*inter,pred+gt),"iou":safe_div(inter,union),"boundary_f1_2px":bf,"hd95_mean_px":float(np.mean(hd)) if hd else float("nan"),"hd95_median_px":float(np.median(hd)) if hd else float("nan")}


def agg_curve(rows,name,errors):
    p="curve_{}_".format(name.lower()); tp=sum(fv(r,p+"tp") for r in rows); fp=sum(fv(r,p+"fp") for r in rows); fn=sum(fv(r,p+"fn") for r in rows); m=f1_counts(tp,fp,fn); ls=sum(fv(r,p+"loc_sum") for r in rows); lc=sum(fv(r,p+"loc_count") for r in rows); hd=finite_values(fv(r,p+"hd95") for r in rows); bp=sum(fv(r,p+"bridge_pred") for r in rows); bt=sum(fv(r,p+"bridge_total") for r in rows); ei=sum(fv(r,p+"centerline_exact_inter") for r in rows); pp=sum(fv(r,p+"pred_points") for r in rows); gp=sum(fv(r,p+"gt_points") for r in rows); return {"presence_precision":m["precision"],"presence_recall":m["recall"],"presence_f1":m["f1"],"mean_localization_error_px":safe_div(ls,lc),"median_localization_error_px":float(np.median(errors)) if len(errors) else float("nan"),"hd95_mean_px":float(np.mean(hd)) if hd else float("nan"),"hd95_median_px":float(np.median(hd)) if hd else float("nan"),"internal_gap_false_bridge_rate":safe_div(bp,bt),"centerline_dice_1px":safe_div(2*ei,pp+gp)}


def agg_hf(rows,e5,e10):
    out={}
    for tag,e in (("5",e5),("10",e10)):
        tp=sum(fv(r,"hf_tp_"+tag) for r in rows); fp=sum(fv(r,"hf_fp_"+tag) for r in rows); fn=sum(fv(r,"hf_fn_"+tag) for r in rows); m=f1_counts(tp,fp,fn); ls=sum(fv(r,"hf_loc_sum_"+tag) for r in rows); lc=sum(fv(r,"hf_loc_count_"+tag) for r in rows); out[tag+"px"]={"precision":m["precision"],"recall":m["recall"],"f1":m["f1"],"mean_localization_error_px":safe_div(ls,lc),"median_localization_error_px":float(np.median(e)) if len(e) else float("nan")}
    out["count_mae"]=float(np.mean([fv(r,"hf_count_abs_error") for r in rows])); return out


def by_patient(rows):
    d={}
    for r in rows: d.setdefault(str(r["patient_id"]),[]).append(r)
    return d


def patient_metrics(rows):
    o={}
    for name in ("SRF","IRF"):
        p="region_{}_".format(name.lower()); inter=sum(fv(r,p+"inter") for r in rows); pred=sum(fv(r,p+"pred") for r in rows); gt=sum(fv(r,p+"gt") for r in rows); un=sum(fv(r,p+"union") for r in rows); o[name+"_dice"]=safe_div(2*inter,pred+gt) if pred+gt>0 else float("nan"); o[name+"_iou"]=safe_div(inter,un) if un>0 else float("nan")
    for name in ("ELM","EZ"):
        p="curve_{}_".format(name.lower()); tp=sum(fv(r,p+"tp") for r in rows); fp=sum(fv(r,p+"fp") for r in rows); fn=sum(fv(r,p+"fn") for r in rows); o[name+"_presence_f1"]=f1_counts(tp,fp,fn)["f1"]; ls=sum(fv(r,p+"loc_sum") for r in rows); lc=sum(fv(r,p+"loc_count") for r in rows); o[name+"_mean_loc_px"]=safe_div(ls,lc) if lc>0 else float("nan")
    for tag in ("5","10"):
        tp=sum(fv(r,"hf_tp_"+tag) for r in rows); fp=sum(fv(r,"hf_fp_"+tag) for r in rows); fn=sum(fv(r,"hf_fn_"+tag) for r in rows); o["HF_f1_{}px".format(tag)]=f1_counts(tp,fp,fn)["f1"]
    o["HF_count_mae"]=float(np.mean([fv(r,"hf_count_abs_error") for r in rows])); return o


def bootstrap(rows):
    g=by_patient(rows); pats=sorted(g)
    if len(pats)!=130: raise RuntimeError("Expected 130 patients, found {}".format(len(pats)))
    pm={p:patient_metrics(g[p]) for p in pats}; names=sorted(next(iter(pm.values()))); rng=np.random.RandomState(BOOTSTRAP_SEED); result={}
    for name in names:
        vals=finite_values(pm[p][name] for p in pats); point=float(np.mean(vals)) if vals else float("nan"); bs=[]
        for _ in range(BOOTSTRAP_REPLICATES):
            samp=[pats[int(i)] for i in rng.randint(0,len(pats),size=len(pats))]; vv=finite_values(pm[p][name] for p in samp)
            if vv: bs.append(float(np.mean(vv)))
        a=np.asarray(bs); result[name]={"patient_macro_mean":point,"ci95_low":float(np.percentile(a,2.5)),"ci95_high":float(np.percentile(a,97.5)),"bootstrap_replicates":len(a)}
    return result


def aggregate_all(protocol):
    rows=[]; payload=[]; ce={"ELM":[],"EZ":[]}; he={"5":[],"10":[]}; folds={}
    for fold in range(1,6):
        d=EVAL_ROOT/"heldout"/("fold_{}".format(fold)); c=load_json(d/"TEST_EVALUATION_COMPLETE.json"); rr=read_rows(Path(c["per_image_csv"])); rows.extend(rr); payload.extend(load_json_gz(Path(c["hf_payload"]))); ar=np.load(Path(c["distance_arrays"])); ce["ELM"].append(ar["curve_elm_errors"]); ce["EZ"].append(ar["curve_ez_errors"]); he["5"].append(ar["hf_5_errors"]); he["10"].append(ar["hf_10_errors"]); folds[str(fold)]={"region":{n:agg_region(rr,n) for n in ("SRF","IRF")},"curve":{n:agg_curve(rr,n,ar["curve_{}_errors".format(n.lower())]) for n in ("ELM","EZ")},"hf":agg_hf(rr,ar["hf_5_errors"],ar["hf_10_errors"])}
    if len(rows)!=EXPECTED_IMAGES or len({r["stem"] for r in rows})!=EXPECTED_IMAGES: raise RuntimeError("Each image was not evaluated exactly once")
    c={k:np.concatenate(v) if v else np.zeros(0,np.float32) for k,v in ce.items()}; h={k:np.concatenate(v) if v else np.zeros(0,np.float32) for k,v in he.items()}; pooled={"region":{n:agg_region(rows,n) for n in ("SRF","IRF")},"curve":{n:agg_curve(rows,n,c[n]) for n in ("ELM","EZ")},"hf":agg_hf(rows,h["5"],h["10"])}; pooled["hf"]["ap50_10x10_boxes"]=ap50(payload); pooled["hf"]["froc"]=froc(payload); ci=bootstrap(rows)
    fields=sorted(rows[0]);
    with PER_IMAGE_ALL_CSV.open("w",encoding="utf-8",newline="") as f:
        w=csv.DictWriter(f,fieldnames=fields); w.writeheader(); w.writerows(rows)
    sr=[]
    for fam in ("region","curve"):
        for target,mm in pooled[fam].items():
            for metric,value in mm.items(): sr.append({"family":fam,"target":target,"metric":metric,"value":value})
    for target,mm in pooled["hf"].items():
        if target=="froc":
            for op,val in mm["operating_points"].items(): sr.append({"family":"hf","target":"FROC","metric":"sensitivity_at_{}_fp_per_image".format(op),"value":val})
        elif isinstance(mm,dict):
            for metric,value in mm.items(): sr.append({"family":"hf","target":target,"metric":metric,"value":value})
        else: sr.append({"family":"hf","target":"HF","metric":target,"value":mm})
    with SUMMARY_CSV.open("w",encoding="utf-8",newline="") as f:
        w=csv.DictWriter(f,fieldnames=["family","target","metric","value"]); w.writeheader(); w.writerows(sr)
    return {"status":"COMPLETE","scientific_status":"FIVEFOLD_HELDOUT_EVALUATION_COMPLETE_NO_POST_TEST_TUNING","evaluation_protocol_sha256":protocol["evaluation_protocol_sha256"],"script_revision":SCRIPT_REVISION,"script_sha256":current_script_sha(),"test_images":EXPECTED_IMAGES,"test_patients":130,"each_image_tested_exactly_once":True,"fold_summaries":folds,"pooled":pooled,"patient_bootstrap_ci":ci,"per_image_csv":str(PER_IMAGE_ALL_CSV),"summary_csv":str(SUMMARY_CSV),"post_test_tuning_allowed":False}


def main():
    print("\n"+"="*110); print("RVO-ME STEP 09: VALIDATION CALIBRATION / HELD-OUT EVALUATION"); print("SCRIPT REVISION: {}".format(SCRIPT_REVISION)); print("="*110); EVAL_ROOT.mkdir(parents=True,exist_ok=True)
    if HELDOUT_COMPLETE_FILE.exists():
        c=load_json(HELDOUT_COMPLETE_FILE); print("\nSTEP 09 already COMPLETE. Protocol SHA: {}".format(c["evaluation_protocol_sha256"])); print("No re-evaluation performed."); return
    print("\n[1/5] Verifying exact model and all five frozen training checkpoints..."); step07,source_sha=load_exact_step07(); manifest=verify_training_manifest(source_sha); images=step07.images_by_stem(); available=sorted(images); device=torch.device("cuda"); print("  Five-fold training manifest: PASS"); print("  Outer tests previously untouched: PASS")
    if not EVAL_PROTOCOL_FILE.exists():
        print("\n[2/5] FIRST RUN — VALIDATION ONLY. TEST roles will NOT be resolved."); cal=run_calibration(manifest,step07,images,available,device); protocol=make_protocol(cal); save_json(EVAL_PROTOCOL_FILE,protocol); append_log_once("<!-- STEP09_EVALUATION_PROTOCOL_FROZEN -->","Step09 evaluation protocol frozen from validation only; no test role resolved. SHA {}.".format(protocol["evaluation_protocol_sha256"])); print("\n"+"="*110); print("STEP 09 CALIBRATION COMPLETE — HELD-OUT TEST STILL UNTOUCHED"); print("="*110); print("Protocol: {}".format(EVAL_PROTOCOL_FILE)); print("SHA: {}".format(protocol["evaluation_protocol_sha256"])); print("\nRun THIS SAME UNCHANGED SCRIPT a second time for the one-time held-out evaluation."); return
    print("\n[2/5] SECOND RUN — verifying frozen evaluation protocol..."); protocol=load_json(EVAL_PROTOCOL_FILE)
    if protocol.get("status")!="FROZEN_BEFORE_HELDOUT_TEST" or protocol.get("script_sha256")!=current_script_sha(): raise RuntimeError("Evaluation protocol/script changed after calibration")
    tmp=dict(protocol); stored=tmp.pop("evaluation_protocol_sha256")
    if stable_sha(tmp)!=stored: raise RuntimeError("Evaluation protocol SHA mismatch")
    print("  Frozen evaluation protocol: PASS"); print("  Script unchanged since calibration: PASS")
    print("\n[3/5] Resolving TEST roles for the first time..."); eye2pat=load_eye_to_patient(); tests={}; seen=set(); seenp=set()
    for fold in range(1,6):
        s=resolve_split_role(fold,"test",available); pats={patient_for_stem(x,eye2pat) for x in s}
        if len(pats)!=EXPECTED_TEST_PATIENTS or seen&set(s) or seenp&pats: raise RuntimeError("Fold {} test leakage/coverage failure".format(fold))
        seen.update(s); seenp.update(pats); tests[fold]=s; print("  Fold {}: {} images / {} patients PASS".format(fold,len(s),len(pats)))
    if len(seen)!=3012 or len(seenp)!=130: raise RuntimeError("Five test folds do not cover all images/patients exactly once")
    start={"status":"STARTED","evaluation_protocol_sha256":stored,"script_sha256":current_script_sha(),"test_images":len(seen),"test_patients":len(seenp),"post_test_tuning_allowed":False}
    if HELDOUT_STARTED_FILE.exists():
        if load_json(HELDOUT_STARTED_FILE).get("evaluation_protocol_sha256")!=stored: raise RuntimeError("Heldout-start protocol mismatch")
    else: save_json(HELDOUT_STARTED_FILE,start)
    print("\n[4/5] Running frozen held-out evaluation...")
    for fold in range(1,6):
        cp=Path(manifest["folds"][str(fold)]["best_checkpoint"]); th=protocol["threshold_calibration"][str(fold)]; res=evaluate_test_fold(fold,tests[fold],cp,th,step07,images,eye2pat,device,stored); print("  Fold {}: {}".format(fold,res["status"]))
    print("\n[5/5] Aggregating metrics and 2,000-patient-bootstrap CIs..."); final=aggregate_all(protocol); save_json(HELDOUT_COMPLETE_FILE,final); append_log_once("<!-- STEP09_HELDOUT_EVALUATION_COMPLETE -->","Step09 five-fold held-out evaluation complete under frozen protocol SHA {}; each image/patient evaluated exactly once; no post-test tuning permitted.".format(stored)); print("\n"+"="*110); print("STEP 09 HELD-OUT EVALUATION COMPLETE"); print("="*110); print("Results: {}".format(HELDOUT_COMPLETE_FILE)); print("Summary CSV: {}".format(SUMMARY_CSV)); print("Per-image CSV: {}".format(PER_IMAGE_ALL_CSV)); print("Post-test tuning allowed: NO")


if __name__ == "__main__":
    main()
