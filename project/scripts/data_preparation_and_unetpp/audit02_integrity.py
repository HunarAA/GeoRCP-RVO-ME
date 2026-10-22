from __future__ import annotations

import csv
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

import cv2
import numpy as np
from openpyxl import load_workbook


# =====================================================================
# CONFIGURATION
# =====================================================================

DATASET_ROOT = Path(r"E:\Papers\RVO_ME\RVO-Lesion")

RAW_LABELME_DIR = DATASET_ROOT / "RVO_Lesion_Labelme"

IMAGE_SEG_DIR = DATASET_ROOT / "Image_Seg"
SEG_IMAGES_DIR = IMAGE_SEG_DIR / "images"
SEG_MASKS_DIR = IMAGE_SEG_DIR / "masks"

TRAIN_TXT = IMAGE_SEG_DIR / "train.txt"
TEST_TXT = IMAGE_SEG_DIR / "test.txt"

OUTPUT_DIR = DATASET_ROOT.parent / "audit02_results"

# Primary targets used by the published segmentation/detection tasks
PRIMARY_LABELS = {"IRF", "SRF", "ELM", "EZ", "HF"}

# Segmentation mask classes used by the official baseline
CLASS_VALUES = {
    "BACKGROUND": 0,
    "SRF": 1,
    "IRF": 2,
    "ELM": 3,
    "EZ": 4,
}

VALUE_TO_CLASS = {
    0: "BACKGROUND",
    1: "SRF",
    2: "IRF",
    3: "ELM",
    4: "EZ",
}

IMAGE_EXTENSIONS = {
    ".jpg",
    ".jpeg",
    ".png",
    ".bmp",
    ".tif",
    ".tiff",
}

# Save difference images only for the worst mask mismatches.
SAVE_TOP_MISMATCH_IMAGES = True
TOP_MISMATCH_COUNT = 20


# =====================================================================
# BASIC HELPERS
# =====================================================================

def safe_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def canonical_id(value: Any) -> str:
    """
    Convert IDs read from Excel into a stable string.

    Examples:
        1       -> "1"
        1.0     -> "1"
        "001"   -> "001"
    """

    if value is None:
        return ""

    if isinstance(value, bool):
        return str(value)

    if isinstance(value, int):
        return str(value)

    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)

    text = str(value).strip()

    # Convert textual "1.0" to "1"
    try:
        numeric = float(text)

        if numeric.is_integer() and "." in text:
            return str(int(numeric))
    except ValueError:
        pass

    return text


def normalize_column_name(value: Any) -> str:
    return re.sub(
        r"[^a-z0-9]",
        "",
        safe_text(value).lower(),
    )


def natural_key(text: str):
    return [
        int(part) if part.isdigit() else part.lower()
        for part in re.split(r"(\d+)", str(text))
    ]


def get_eye_id_from_filename(filename: str) -> Optional[str]:
    """
    Example:
        72_10.jpg -> 72

    The RVO-ME paper states that the filename prefix is Eye ID.
    """

    stem = Path(filename).stem

    match = re.match(r"^([^_]+)_", stem)

    if match:
        return canonical_id(match.group(1))

    return None


def read_split_file(path: Path) -> List[str]:

    if not path.exists():
        raise FileNotFoundError(
            f"Split file not found:\n{path}"
        )

    values = []

    with path.open(
        "r",
        encoding="utf-8-sig",
    ) as f:

        for line in f:
            line = line.strip()

            if line:
                values.append(
                    Path(line).name
                )

    return values


def write_csv(
    path: Path,
    rows: List[Dict[str, Any]],
    fieldnames: Optional[List[str]] = None,
):

    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    if not rows:
        path.write_text(
            "",
            encoding="utf-8",
        )
        return

    if fieldnames is None:
        fieldnames = list(
            rows[0].keys()
        )

    with path.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as f:

        writer = csv.DictWriter(
            f,
            fieldnames=fieldnames,
        )

        writer.writeheader()
        writer.writerows(rows)


def save_json(path: Path, data: Any):

    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with path.open(
        "w",
        encoding="utf-8",
    ) as f:

        json.dump(
            data,
            f,
            indent=2,
            ensure_ascii=False,
            default=str,
        )


# =====================================================================
# FIND FILES
# =====================================================================

def find_demographics_file() -> Path:

    candidates = list(
        DATASET_ROOT.glob("*.xlsx")
    )

    if not candidates:
        raise FileNotFoundError(
            "No demographics XLSX file found."
        )

    for path in candidates:

        if "demograph" in path.name.lower():
            return path

    return candidates[0]


def build_raw_file_lookup():

    print("Building raw LabelMe file lookup...")

    json_lookup = {}
    image_lookup = {}

    for path in RAW_LABELME_DIR.rglob("*"):

        if not path.is_file():
            continue

        stem = path.stem

        if path.suffix.lower() == ".json":

            if stem in json_lookup:
                raise RuntimeError(
                    f"Duplicate JSON stem found: {stem}"
                )

            json_lookup[stem] = path

        elif path.suffix.lower() in IMAGE_EXTENSIONS:

            if stem in image_lookup:
                raise RuntimeError(
                    f"Duplicate image stem found: {stem}"
                )

            image_lookup[stem] = path

    return json_lookup, image_lookup


def build_segmentation_mask_lookup():

    mask_lookup = {}

    for path in SEG_MASKS_DIR.iterdir():

        if (
            path.is_file()
            and path.suffix.lower()
            in IMAGE_EXTENSIONS
        ):

            if path.stem in mask_lookup:
                raise RuntimeError(
                    "Duplicate segmentation mask "
                    f"stem: {path.stem}"
                )

            mask_lookup[path.stem] = path

    return mask_lookup


# =====================================================================
# DEMOGRAPHICS
# =====================================================================

def load_demographics():

    print("\n[1/4] Loading demographics...")

    demographics_path = (
        find_demographics_file()
    )

    workbook = load_workbook(
        demographics_path,
        read_only=True,
        data_only=True,
    )

    sheet = workbook.active

    all_rows = list(
        sheet.iter_rows(
            values_only=True
        )
    )

    if not all_rows:
        raise RuntimeError(
            "Demographics Excel file is empty."
        )

    headers = [
        safe_text(v)
        for v in all_rows[0]
    ]

    normalized_headers = [
        normalize_column_name(v)
        for v in headers
    ]

    column_indices = {}

    for index, column in enumerate(
        normalized_headers
    ):

        if column in {
            "patientid",
            "patient",
        }:
            column_indices["patient"] = index

        elif column in {
            "eyeid",
            "eye",
        }:
            column_indices["eye"] = index

        elif column.startswith("age"):
            column_indices["age"] = index

        elif column in {
            "gender",
            "sex",
        }:
            column_indices["gender"] = index

        elif (
            "eyecategory" in column
            or "eyecateg" in column
            or column == "laterality"
        ):
            column_indices[
                "eye_category"
            ] = index

    required = [
        "patient",
        "eye",
    ]

    for field in required:

        if field not in column_indices:
            raise RuntimeError(
                f"Could not identify {field} "
                "column in demographics file."
            )

    rows = []

    eye_to_patient = {}

    for excel_row, values in enumerate(
        all_rows[1:],
        start=2,
    ):

        if not any(
            value is not None
            for value in values
        ):
            continue

        def value_at(field):
            idx = column_indices.get(field)

            if idx is None:
                return None

            if idx >= len(values):
                return None

            return values[idx]

        patient_id = canonical_id(
            value_at("patient")
        )

        eye_id = canonical_id(
            value_at("eye")
        )

        record = {
            "excel_row": excel_row,
            "patient_id": patient_id,
            "eye_id": eye_id,
            "eye_category":
                safe_text(
                    value_at("eye_category")
                ),
            "age": value_at("age"),
            "gender":
                safe_text(
                    value_at("gender")
                ),
        }

        rows.append(record)

        if eye_id:

            if (
                eye_id in eye_to_patient
                and eye_to_patient[eye_id]
                != patient_id
            ):
                raise RuntimeError(
                    "One Eye ID maps to multiple "
                    f"patients: Eye {eye_id}"
                )

            eye_to_patient[
                eye_id
            ] = patient_id

    workbook.close()

    return {
        "path": demographics_path,
        "rows": rows,
        "eye_to_patient":
            eye_to_patient,
    }


# =====================================================================
# PATIENT-LEVEL DEMOGRAPHIC AUDIT
# =====================================================================

def audit_patient_demographics(
    demographics
):

    patient_values = defaultdict(
        lambda: {
            "eyes": set(),
            "ages": set(),
            "genders": set(),
            "eye_categories": set(),
        }
    )

    for row in demographics["rows"]:

        patient = row[
            "patient_id"
        ]

        if not patient:
            continue

        patient_values[
            patient
        ]["eyes"].add(
            row["eye_id"]
        )

        if row["age"] is not None:
            patient_values[
                patient
            ]["ages"].add(
                safe_text(
                    row["age"]
                )
            )

        if row["gender"]:
            patient_values[
                patient
            ]["genders"].add(
                row["gender"]
            )

        if row["eye_category"]:
            patient_values[
                patient
            ][
                "eye_categories"
            ].add(
                row[
                    "eye_category"
                ]
            )

    patient_rows = []
    inconsistencies = []

    for patient in sorted(
        patient_values,
        key=natural_key,
    ):

        info = patient_values[
            patient
        ]

        ages = sorted(
            info["ages"]
        )

        genders = sorted(
            info["genders"]
        )

        eyes = sorted(
            info["eyes"],
            key=natural_key,
        )

        patient_rows.append({
            "patient_id": patient,
            "number_of_eyes":
                len(eyes),
            "eye_ids":
                ";".join(eyes),
            "ages":
                ";".join(ages),
            "genders":
                ";".join(genders),
            "eye_categories":
                ";".join(
                    sorted(
                        info[
                            "eye_categories"
                        ]
                    )
                ),
        })

        reasons = []

        if len(ages) > 1:
            reasons.append(
                "multiple_age_values"
            )

        if len(genders) > 1:
            reasons.append(
                "multiple_gender_values"
            )

        if reasons:

            inconsistencies.append({
                "patient_id": patient,
                "reasons":
                    ";".join(reasons),
                "eye_ids":
                    ";".join(eyes),
                "ages":
                    ";".join(ages),
                "genders":
                    ";".join(genders),
            })

    patient_gender_counts = Counter()

    for row in patient_rows:

        genders = [
            g
            for g in
            row["genders"].split(";")
            if g
        ]

        if len(genders) == 1:
            patient_gender_counts[
                genders[0]
            ] += 1
        else:
            patient_gender_counts[
                "__AMBIGUOUS__"
            ] += 1

    return {
        "patient_rows": patient_rows,
        "inconsistencies":
            inconsistencies,
        "patient_gender_counts":
            patient_gender_counts,
    }


# =====================================================================
# SPLIT INTEGRITY
# =====================================================================

def audit_split_integrity(
    demographics
):

    print(
        "[2/4] Auditing train/test "
        "patient integrity..."
    )

    train_names = read_split_file(
        TRAIN_TXT
    )

    test_names = read_split_file(
        TEST_TXT
    )

    eye_to_patient = demographics[
        "eye_to_patient"
    ]

    split_records = []

    train_eyes = set()
    test_eyes = set()

    train_patients = set()
    test_patients = set()

    per_eye = defaultdict(
        lambda: {
            "train": [],
            "test": [],
        }
    )

    per_patient = defaultdict(
        lambda: {
            "train_files": [],
            "test_files": [],
            "train_eyes": set(),
            "test_eyes": set(),
        }
    )

    for split_name, filenames in [
        ("train", train_names),
        ("test", test_names),
    ]:

        for filename in filenames:

            eye_id = (
                get_eye_id_from_filename(
                    filename
                )
            )

            if eye_id is None:

                patient_id = ""

            else:

                patient_id = (
                    eye_to_patient.get(
                        eye_id,
                        "",
                    )
                )

            record = {
                "split": split_name,
                "filename": filename,
                "eye_id":
                    eye_id or "",
                "patient_id":
                    patient_id,
                "eye_found_in_demographics":
                    bool(patient_id),
            }

            split_records.append(
                record
            )

            if not eye_id:
                continue

            per_eye[
                eye_id
            ][split_name].append(
                filename
            )

            if split_name == "train":
                train_eyes.add(
                    eye_id
                )
            else:
                test_eyes.add(
                    eye_id
                )

            if patient_id:

                if split_name == "train":
                    train_patients.add(
                        patient_id
                    )
                else:
                    test_patients.add(
                        patient_id
                    )

                per_patient[
                    patient_id
                ][
                    f"{split_name}_files"
                ].append(
                    filename
                )

                per_patient[
                    patient_id
                ][
                    f"{split_name}_eyes"
                ].add(
                    eye_id
                )

    overlapping_eyes = sorted(
        train_eyes & test_eyes,
        key=natural_key,
    )

    overlapping_patients = sorted(
        train_patients & test_patients,
        key=natural_key,
    )

    eye_overlap_rows = []

    for eye_id in overlapping_eyes:

        train_files = sorted(
            per_eye[
                eye_id
            ]["train"],
            key=natural_key,
        )

        test_files = sorted(
            per_eye[
                eye_id
            ]["test"],
            key=natural_key,
        )

        eye_overlap_rows.append({
            "eye_id": eye_id,
            "patient_id":
                eye_to_patient.get(
                    eye_id,
                    "",
                ),
            "train_image_count":
                len(train_files),
            "test_image_count":
                len(test_files),
            "train_files":
                ";".join(
                    train_files
                ),
            "test_files":
                ";".join(
                    test_files
                ),
        })

    patient_overlap_rows = []

    for patient_id in overlapping_patients:

        info = per_patient[
            patient_id
        ]

        train_files = sorted(
            info["train_files"],
            key=natural_key,
        )

        test_files = sorted(
            info["test_files"],
            key=natural_key,
        )

        train_patient_eyes = sorted(
            info["train_eyes"],
            key=natural_key,
        )

        test_patient_eyes = sorted(
            info["test_eyes"],
            key=natural_key,
        )

        shared_eyes = sorted(
            set(train_patient_eyes)
            & set(test_patient_eyes),
            key=natural_key,
        )

        patient_overlap_rows.append({
            "patient_id":
                patient_id,
            "train_eye_ids":
                ";".join(
                    train_patient_eyes
                ),
            "test_eye_ids":
                ";".join(
                    test_patient_eyes
                ),
            "shared_eye_ids":
                ";".join(
                    shared_eyes
                ),
            "train_image_count":
                len(train_files),
            "test_image_count":
                len(test_files),
            "train_files":
                ";".join(
                    train_files
                ),
            "test_files":
                ";".join(
                    test_files
                ),
        })

    # -------------------------------------------------------------
    # Patient assignment table for all patients
    # -------------------------------------------------------------

    patient_assignment_rows = []

    all_patients = sorted(
        train_patients | test_patients,
        key=natural_key,
    )

    for patient_id in all_patients:

        info = per_patient[
            patient_id
        ]

        in_train = (
            patient_id
            in train_patients
        )

        in_test = (
            patient_id
            in test_patients
        )

        if in_train and in_test:
            status = "BOTH"
        elif in_train:
            status = "TRAIN_ONLY"
        else:
            status = "TEST_ONLY"

        patient_assignment_rows.append({
            "patient_id":
                patient_id,
            "status":
                status,
            "train_eye_ids":
                ";".join(
                    sorted(
                        info[
                            "train_eyes"
                        ],
                        key=natural_key,
                    )
                ),
            "test_eye_ids":
                ";".join(
                    sorted(
                        info[
                            "test_eyes"
                        ],
                        key=natural_key,
                    )
                ),
            "train_image_count":
                len(
                    info[
                        "train_files"
                    ]
                ),
            "test_image_count":
                len(
                    info[
                        "test_files"
                    ]
                ),
        })

    unmapped = [
        row
        for row in split_records
        if not row[
            "eye_found_in_demographics"
        ]
    ]

    return {
        "train_names":
            train_names,
        "test_names":
            test_names,

        "train_eyes":
            train_eyes,
        "test_eyes":
            test_eyes,

        "train_patients":
            train_patients,
        "test_patients":
            test_patients,

        "overlapping_eyes":
            overlapping_eyes,
        "overlapping_patients":
            overlapping_patients,

        "eye_overlap_rows":
            eye_overlap_rows,

        "patient_overlap_rows":
            patient_overlap_rows,

        "patient_assignment_rows":
            patient_assignment_rows,

        "split_records":
            split_records,

        "unmapped":
            unmapped,
    }


# =====================================================================
# ANNOTATION ANOMALY AUDIT
# =====================================================================

def classify_annotation_status(
    label: str,
    shape_type: str,
    n_points: int,
) -> Tuple[str, str]:
    """
    Determine whether an annotation is normal for our intended
    geometry-aware formulation.

    IRF / SRF:
        Region contours. Both polygon and linestrip appear in the
        release and the official conversion code fills them.

    ELM / EZ:
        Expected as linestrip.

    HF:
        Expected as point.

    Other labels:
        Retained for audit but not primary targets.
    """

    if label in {"IRF", "SRF"}:

        if shape_type not in {
            "polygon",
            "linestrip",
        }:
            return (
                "UNUSUAL",
                "region_target_with_"
                f"{shape_type}",
            )

        if n_points < 3:
            return (
                "UNUSUAL",
                "region_target_has_"
                "fewer_than_3_points",
            )

        return (
            "EXPECTED",
            "region_contour",
        )

    if label in {"ELM", "EZ"}:

        if shape_type != "linestrip":
            return (
                "UNUSUAL",
                "curve_target_with_"
                f"{shape_type}",
            )

        if n_points < 2:
            return (
                "UNUSUAL",
                "curve_target_has_"
                "fewer_than_2_points",
            )

        return (
            "EXPECTED",
            "curve_linestrip",
        )

    if label == "HF":

        if shape_type != "point":
            return (
                "UNUSUAL",
                "point_target_with_"
                f"{shape_type}",
            )

        if n_points != 1:
            return (
                "UNUSUAL",
                "point_target_has_"
                f"{n_points}_points",
            )

        return (
            "EXPECTED",
            "point",
        )

    return (
        "NON_PRIMARY",
        "non_primary_label",
    )


def audit_annotations(
    json_lookup
):

    print(
        "[3/4] Auditing unusual "
        "raw annotations..."
    )

    all_rows = []
    unusual_rows = []

    label_shape_counts = Counter()
    status_counts = Counter()

    label_counts = Counter()

    for index, stem in enumerate(
        sorted(
            json_lookup,
            key=natural_key,
        ),
        start=1,
    ):

        json_path = (
            json_lookup[stem]
        )

        with json_path.open(
            "r",
            encoding="utf-8-sig",
        ) as f:

            data = json.load(f)

        shapes = data.get(
            "shapes",
            [],
        )

        for shape_index, shape in enumerate(
            shapes
        ):

            label = safe_text(
                shape.get("label")
            )

            shape_type = (
                safe_text(
                    shape.get(
                        "shape_type"
                    )
                )
                or "polygon"
            )

            points = shape.get(
                "points",
                [],
            )

            if not isinstance(
                points,
                list,
            ):
                points = []

            n_points = len(
                points
            )

            status, reason = (
                classify_annotation_status(
                    label,
                    shape_type,
                    n_points,
                )
            )

            eye_id = (
                get_eye_id_from_filename(
                    stem
                )
            )

            row = {
                "image_stem":
                    stem,
                "eye_id":
                    eye_id or "",
                "json_path":
                    str(json_path),
                "shape_index":
                    shape_index,
                "label":
                    label,
                "shape_type":
                    shape_type,
                "n_points":
                    n_points,
                "status":
                    status,
                "reason":
                    reason,
                "points_json":
                    json.dumps(
                        points,
                        ensure_ascii=False,
                    ),
            }

            all_rows.append(
                row
            )

            label_counts[
                label
            ] += 1

            label_shape_counts[
                (
                    label,
                    shape_type,
                )
            ] += 1

            status_counts[
                status
            ] += 1

            if status != "EXPECTED":
                unusual_rows.append(
                    row
                )

        if index % 500 == 0:
            print(
                f"    Parsed "
                f"{index}/"
                f"{len(json_lookup)} JSON files"
            )

    return {
        "all_rows":
            all_rows,
        "unusual_rows":
            unusual_rows,
        "label_counts":
            label_counts,
        "label_shape_counts":
            label_shape_counts,
        "status_counts":
            status_counts,
    }


# =====================================================================
# MASK RECONSTRUCTION
# =====================================================================

def load_json(
    json_path: Path
):

    with json_path.open(
        "r",
        encoding="utf-8-sig",
    ) as f:

        return json.load(f)


def reconstruct_mask(
    json_path: Path,
    image_path: Path,
    target_only: bool,
) -> np.ndarray:
    """
    Reconstruct a segmentation mask following the logic of the
    authors' official convert-json-to-masks.py.

    target_only=False:
        Reproduce repository behaviour as closely as possible,
        including drawing non-target polygon/point labels with
        value 0.

    target_only=True:
        Ignore non-target annotations entirely. This lets us test
        whether those zero-valued annotations affect published masks.
    """

    image = cv2.imread(
        str(image_path),
        cv2.IMREAD_COLOR,
    )

    if image is None:
        raise RuntimeError(
            f"Cannot read raw image:\n{image_path}"
        )

    # Official code starts with a 3-channel uint8 mask.
    mask = np.zeros(
        image.shape,
        dtype=np.uint8,
    )

    data = load_json(
        json_path
    )

    shapes = data.get(
        "shapes",
        [],
    )

    for shape in shapes:

        label = safe_text(
            shape.get("label")
        )

        if (
            target_only
            and label
            not in {
                "SRF",
                "IRF",
                "ELM",
                "EZ",
            }
        ):
            continue

        points = shape.get(
            "points",
            [],
        )

        shape_type = (
            safe_text(
                shape.get(
                    "shape_type"
                )
            )
            or "polygon"
        )

        if not isinstance(
            points,
            list,
        ):
            continue

        if len(points) == 0:
            continue

        # Official class encoding
        if label == "SRF":
            color = (1, 1, 1)

        elif label == "IRF":
            color = (2, 2, 2)

        elif label == "ELM":
            color = (3, 3, 3)

        elif label == "EZ":
            color = (4, 4, 4)

        else:
            # Exact behaviour of official conversion script:
            # unrecognised labels use zero.
            color = (0, 0, 0)

        try:

            pts = np.array(
                points,
                dtype=np.int32,
            )

            pts = pts.reshape(
                (-1, 1, 2)
            )

        except Exception:
            continue

        # ---------------------------------------------------------
        # Reproduce official repository decision logic
        # ---------------------------------------------------------

        if (
            shape_type == "polygon"
            or label in {
                "IRF",
                "SRF",
            }
        ):

            if len(pts) < 3:
                continue

            cv2.fillPoly(
                mask,
                [pts],
                color,
            )

        elif (
            shape_type == "linestrip"
            and label
            in {
                "ELM",
                "EZ",
            }
        ):

            if len(pts) < 2:
                continue

            cv2.polylines(
                mask,
                [pts],
                isClosed=False,
                color=color,
                thickness=2,
            )

        elif shape_type == "point":

            center = tuple(
                pts[0][0]
            )

            cv2.circle(
                mask,
                center,
                radius=2,
                color=color,
                thickness=-1,
            )

        # All other cases are ignored,
        # matching the repository's unsupported-shape branch.

    mask_gray = cv2.cvtColor(
        mask,
        cv2.COLOR_BGR2GRAY,
    )

    return mask_gray


# =====================================================================
# MASK COMPARISON METRICS
# =====================================================================

def compare_masks(
    reference: np.ndarray,
    reconstructed: np.ndarray,
):

    if reference.shape != reconstructed.shape:

        return {
            "same_shape": False,
            "exact_match": False,
            "different_pixels": None,
            "difference_percent": None,
            "max_absolute_difference": None,
        }

    difference = cv2.absdiff(
        reference,
        reconstructed,
    )

    different_pixels = int(
        np.count_nonzero(
            difference
        )
    )

    total_pixels = int(
        reference.size
    )

    return {
        "same_shape": True,

        "exact_match":
            different_pixels == 0,

        "different_pixels":
            different_pixels,

        "difference_percent":
            (
                100.0
                * different_pixels
                / total_pixels
            ),

        "max_absolute_difference":
            int(
                difference.max()
            ),
    }


def update_confusion_matrix(
    matrix: np.ndarray,
    reference: np.ndarray,
    prediction: np.ndarray,
):

    if (
        reference.shape
        != prediction.shape
    ):
        return

    reference_flat = (
        reference.ravel()
    )

    prediction_flat = (
        prediction.ravel()
    )

    valid = (
        (reference_flat >= 0)
        & (reference_flat <= 4)
        & (prediction_flat >= 0)
        & (prediction_flat <= 4)
    )

    encoded = (
        reference_flat[valid].astype(
            np.int64
        )
        * 5
        + prediction_flat[valid].astype(
            np.int64
        )
    )

    counts = np.bincount(
        encoded,
        minlength=25,
    ).reshape(
        5,
        5,
    )

    matrix += counts


def confusion_matrix_to_rows(
    matrix: np.ndarray
):

    rows = []

    for reference_value in range(5):

        row = {
            "reference_class":
                VALUE_TO_CLASS[
                    reference_value
                ],
            "reference_value":
                reference_value,
        }

        for predicted_value in range(5):

            key = (
                "pred_"
                + VALUE_TO_CLASS[
                    predicted_value
                ]
            )

            row[key] = int(
                matrix[
                    reference_value,
                    predicted_value,
                ]
            )

        rows.append(
            row
        )

    return rows


def class_agreement_from_confusion(
    matrix: np.ndarray
):

    rows = []

    for class_value in range(1, 5):

        tp = int(
            matrix[
                class_value,
                class_value,
            ]
        )

        fp = int(
            matrix[
                :,
                class_value,
            ].sum()
            - tp
        )

        fn = int(
            matrix[
                class_value,
                :,
            ].sum()
            - tp
        )

        denominator_dice = (
            2 * tp
            + fp
            + fn
        )

        denominator_iou = (
            tp
            + fp
            + fn
        )

        dice = (
            (2 * tp)
            / denominator_dice
            if denominator_dice
            else 1.0
        )

        iou = (
            tp
            / denominator_iou
            if denominator_iou
            else 1.0
        )

        rows.append({
            "class_value":
                class_value,
            "class_name":
                VALUE_TO_CLASS[
                    class_value
                ],
            "true_positive":
                tp,
            "false_positive":
                fp,
            "false_negative":
                fn,
            "dice":
                dice,
            "iou":
                iou,
        })

    return rows


# =====================================================================
# DIFFERENCE VISUALIZATION
# =====================================================================

def save_difference_visualization(
    output_path: Path,
    reference: np.ndarray,
    reconstructed: np.ndarray,
):

    difference = (
        reference
        != reconstructed
    ).astype(
        np.uint8
    ) * 255

    # Scale class values to visible grayscale.
    reference_visible = (
        reference.astype(
            np.uint8
        )
        * 60
    )

    reconstructed_visible = (
        reconstructed.astype(
            np.uint8
        )
        * 60
    )

    combined = np.hstack([
        reference_visible,
        reconstructed_visible,
        difference,
    ])

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    cv2.imwrite(
        str(output_path),
        combined,
    )


# =====================================================================
# RECONSTRUCT AND COMPARE ALL MASKS
# =====================================================================

def audit_mask_reconstruction(
    json_lookup,
    image_lookup,
):

    print(
        "[4/4] Reconstructing all "
        "3012 segmentation masks..."
    )

    mask_lookup = (
        build_segmentation_mask_lookup()
    )

    comparison_rows = []

    missing_rows = []

    confusion_repo_exact = np.zeros(
        (5, 5),
        dtype=np.int64,
    )

    confusion_target_only = np.zeros(
        (5, 5),
        dtype=np.int64,
    )

    unexpected_reference_values = (
        Counter()
    )

    stems = sorted(
        mask_lookup,
        key=natural_key,
    )

    for index, stem in enumerate(
        stems,
        start=1,
    ):

        official_path = (
            mask_lookup[stem]
        )

        json_path = (
            json_lookup.get(stem)
        )

        image_path = (
            image_lookup.get(stem)
        )

        if (
            json_path is None
            or image_path is None
        ):

            missing_rows.append({
                "image_stem":
                    stem,
                "json_found":
                    json_path is not None,
                "raw_image_found":
                    image_path is not None,
                "official_mask_path":
                    str(
                        official_path
                    ),
            })

            continue

        official_mask = cv2.imread(
            str(official_path),
            cv2.IMREAD_GRAYSCALE,
        )

        if official_mask is None:

            missing_rows.append({
                "image_stem":
                    stem,
                "json_found":
                    True,
                "raw_image_found":
                    True,
                "official_mask_path":
                    str(
                        official_path
                    ),
                "problem":
                    "cannot_read_official_mask",
            })

            continue

        for value, count in zip(
            *np.unique(
                official_mask,
                return_counts=True,
            )
        ):

            if int(value) not in range(5):

                unexpected_reference_values[
                    int(value)
                ] += int(count)

        # ---------------------------------------------------------
        # Reconstruction 1:
        # Match official GitHub conversion behaviour
        # ---------------------------------------------------------

        repo_exact_mask = (
            reconstruct_mask(
                json_path,
                image_path,
                target_only=False,
            )
        )

        # ---------------------------------------------------------
        # Reconstruction 2:
        # Ignore ERM/HF/PED/HRD/etc. during segmentation conversion
        # ---------------------------------------------------------

        target_only_mask = (
            reconstruct_mask(
                json_path,
                image_path,
                target_only=True,
            )
        )

        comparison_exact = (
            compare_masks(
                official_mask,
                repo_exact_mask,
            )
        )

        comparison_target = (
            compare_masks(
                official_mask,
                target_only_mask,
            )
        )

        update_confusion_matrix(
            confusion_repo_exact,
            official_mask,
            repo_exact_mask,
        )

        update_confusion_matrix(
            confusion_target_only,
            official_mask,
            target_only_mask,
        )

        comparison_rows.append({
            "image_stem":
                stem,

            "official_mask_path":
                str(
                    official_path
                ),

            "repo_exact_match":
                comparison_exact[
                    "exact_match"
                ],

            "repo_exact_different_pixels":
                comparison_exact[
                    "different_pixels"
                ],

            "repo_exact_difference_percent":
                comparison_exact[
                    "difference_percent"
                ],

            "repo_exact_max_abs_difference":
                comparison_exact[
                    "max_absolute_difference"
                ],

            "target_only_exact_match":
                comparison_target[
                    "exact_match"
                ],

            "target_only_different_pixels":
                comparison_target[
                    "different_pixels"
                ],

            "target_only_difference_percent":
                comparison_target[
                    "difference_percent"
                ],

            "target_only_max_abs_difference":
                comparison_target[
                    "max_absolute_difference"
                ],
        })

        if index % 250 == 0:

            print(
                f"    Compared "
                f"{index}/"
                f"{len(stems)} masks"
            )

    # -------------------------------------------------------------
    # Save visualizations for worst mismatches
    # -------------------------------------------------------------

    if SAVE_TOP_MISMATCH_IMAGES:

        mismatches = [
            row
            for row in comparison_rows
            if (
                row[
                    "repo_exact_different_pixels"
                ]
                not in {
                    None,
                    0,
                }
            )
        ]

        mismatches.sort(
            key=lambda x:
                x[
                    "repo_exact_different_pixels"
                ],
            reverse=True,
        )

        visual_dir = (
            OUTPUT_DIR
            / "top_mask_differences"
        )

        for rank, row in enumerate(
            mismatches[
                :TOP_MISMATCH_COUNT
            ],
            start=1,
        ):

            stem = row[
                "image_stem"
            ]

            official_mask = (
                cv2.imread(
                    str(
                        mask_lookup[
                            stem
                        ]
                    ),
                    cv2.IMREAD_GRAYSCALE,
                )
            )

            reconstructed = (
                reconstruct_mask(
                    json_lookup[
                        stem
                    ],
                    image_lookup[
                        stem
                    ],
                    target_only=False,
                )
            )

            save_difference_visualization(
                visual_dir
                / (
                    f"{rank:02d}_"
                    f"{stem}.png"
                ),
                official_mask,
                reconstructed,
            )

    return {
        "comparison_rows":
            comparison_rows,

        "missing_rows":
            missing_rows,

        "confusion_repo_exact":
            confusion_repo_exact,

        "confusion_target_only":
            confusion_target_only,

        "unexpected_reference_values":
            unexpected_reference_values,
    }


# =====================================================================
# REPORTING
# =====================================================================

def save_all_reports(
    demographics,
    patient_demo,
    split,
    annotations,
    mask_audit,
):

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    # -------------------------------------------------------------
    # Demographics
    # -------------------------------------------------------------

    write_csv(
        OUTPUT_DIR
        / "patient_demographics.csv",
        patient_demo[
            "patient_rows"
        ],
    )

    write_csv(
        OUTPUT_DIR
        / "patient_demographic_inconsistencies.csv",
        patient_demo[
            "inconsistencies"
        ],
    )

    # -------------------------------------------------------------
    # Split
    # -------------------------------------------------------------

    write_csv(
        OUTPUT_DIR
        / "split_file_mapping.csv",
        split[
            "split_records"
        ],
    )

    write_csv(
        OUTPUT_DIR
        / "split_eye_overlap.csv",
        split[
            "eye_overlap_rows"
        ],
    )

    write_csv(
        OUTPUT_DIR
        / "split_patient_overlap.csv",
        split[
            "patient_overlap_rows"
        ],
    )

    write_csv(
        OUTPUT_DIR
        / "split_patient_assignment.csv",
        split[
            "patient_assignment_rows"
        ],
    )

    write_csv(
        OUTPUT_DIR
        / "split_unmapped_eye_ids.csv",
        split[
            "unmapped"
        ],
    )

    # -------------------------------------------------------------
    # Annotation anomalies
    # -------------------------------------------------------------

    write_csv(
        OUTPUT_DIR
        / "unusual_annotations.csv",
        annotations[
            "unusual_rows"
        ],
    )

    write_csv(
        OUTPUT_DIR
        / "all_annotation_integrity.csv",
        annotations[
            "all_rows"
        ],
    )

    label_shape_rows = []

    for (
        label,
        shape_type,
    ), count in sorted(
        annotations[
            "label_shape_counts"
        ].items()
    ):

        label_shape_rows.append({
            "label":
                label,
            "shape_type":
                shape_type,
            "count":
                count,
        })

    write_csv(
        OUTPUT_DIR
        / "annotation_label_shape_counts.csv",
        label_shape_rows,
    )

    # -------------------------------------------------------------
    # Mask comparison
    # -------------------------------------------------------------

    write_csv(
        OUTPUT_DIR
        / "mask_reconstruction_comparison.csv",
        mask_audit[
            "comparison_rows"
        ],
    )

    write_csv(
        OUTPUT_DIR
        / "mask_reconstruction_missing.csv",
        mask_audit[
            "missing_rows"
        ],
    )

    repo_confusion_rows = (
        confusion_matrix_to_rows(
            mask_audit[
                "confusion_repo_exact"
            ]
        )
    )

    target_confusion_rows = (
        confusion_matrix_to_rows(
            mask_audit[
                "confusion_target_only"
            ]
        )
    )

    write_csv(
        OUTPUT_DIR
        / "mask_confusion_repo_exact.csv",
        repo_confusion_rows,
    )

    write_csv(
        OUTPUT_DIR
        / "mask_confusion_target_only.csv",
        target_confusion_rows,
    )

    write_csv(
        OUTPUT_DIR
        / "mask_class_agreement_repo_exact.csv",
        class_agreement_from_confusion(
            mask_audit[
                "confusion_repo_exact"
            ]
        ),
    )

    write_csv(
        OUTPUT_DIR
        / "mask_class_agreement_target_only.csv",
        class_agreement_from_confusion(
            mask_audit[
                "confusion_target_only"
            ]
        ),
    )

    # -------------------------------------------------------------
    # Summary JSON
    # -------------------------------------------------------------

    repo_rows = mask_audit[
        "comparison_rows"
    ]

    repo_exact_matches = sum(
        1
        for row in repo_rows
        if row[
            "repo_exact_match"
        ]
    )

    target_exact_matches = sum(
        1
        for row in repo_rows
        if row[
            "target_only_exact_match"
        ]
    )

    total_compared = len(
        repo_rows
    )

    summary = {
        "dataset_root":
            str(DATASET_ROOT),

        "demographics": {
            "excel_file":
                str(
                    demographics[
                        "path"
                    ]
                ),

            "eye_rows":
                len(
                    demographics[
                        "rows"
                    ]
                ),

            "unique_eye_ids":
                len(
                    demographics[
                        "eye_to_patient"
                    ]
                ),

            "unique_patients":
                len(
                    patient_demo[
                        "patient_rows"
                    ]
                ),

            "patient_gender_counts":
                dict(
                    patient_demo[
                        "patient_gender_counts"
                    ]
                ),

            "patient_demographic_inconsistencies":
                len(
                    patient_demo[
                        "inconsistencies"
                    ]
                ),
        },

        "official_split_observed": {

            "train_images":
                len(
                    split[
                        "train_names"
                    ]
                ),

            "test_images":
                len(
                    split[
                        "test_names"
                    ]
                ),

            "train_unique_eyes":
                len(
                    split[
                        "train_eyes"
                    ]
                ),

            "test_unique_eyes":
                len(
                    split[
                        "test_eyes"
                    ]
                ),

            "eye_overlap_count":
                len(
                    split[
                        "overlapping_eyes"
                    ]
                ),

            "overlapping_eye_ids":
                split[
                    "overlapping_eyes"
                ],

            "train_unique_patients":
                len(
                    split[
                        "train_patients"
                    ]
                ),

            "test_unique_patients":
                len(
                    split[
                        "test_patients"
                    ]
                ),

            "patient_overlap_count":
                len(
                    split[
                        "overlapping_patients"
                    ]
                ),

            "overlapping_patient_ids":
                split[
                    "overlapping_patients"
                ],

            "unmapped_split_files":
                len(
                    split[
                        "unmapped"
                    ]
                ),
        },

        "paper_reported_split": {
            "training_patients":
                104,
            "test_patients":
                26,
            "patient_overlap_expected":
                0,
        },

        "annotations": {

            "total_annotations":
                len(
                    annotations[
                        "all_rows"
                    ]
                ),

            "status_counts":
                dict(
                    annotations[
                        "status_counts"
                    ]
                ),

            "label_counts":
                dict(
                    annotations[
                        "label_counts"
                    ]
                ),

            "unusual_or_non_primary_count":
                len(
                    annotations[
                        "unusual_rows"
                    ]
                ),
        },

        "mask_reconstruction": {

            "masks_compared":
                total_compared,

            "repo_exact_match_count":
                repo_exact_matches,

            "repo_exact_mismatch_count":
                total_compared
                - repo_exact_matches,

            "repo_exact_match_percent":
                (
                    100.0
                    * repo_exact_matches
                    / total_compared
                    if total_compared
                    else 0
                ),

            "target_only_exact_match_count":
                target_exact_matches,

            "target_only_mismatch_count":
                total_compared
                - target_exact_matches,

            "target_only_match_percent":
                (
                    100.0
                    * target_exact_matches
                    / total_compared
                    if total_compared
                    else 0
                ),

            "missing_inputs":
                len(
                    mask_audit[
                        "missing_rows"
                    ]
                ),

            "unexpected_official_mask_values":
                dict(
                    mask_audit[
                        "unexpected_reference_values"
                    ]
                ),
        },
    }

    save_json(
        OUTPUT_DIR
        / "audit02_summary.json",
        summary,
    )

    return summary


# =====================================================================
# TERMINAL SUMMARY
# =====================================================================

def print_summary(
    summary,
    split,
    annotations,
):

    print("\n")
    print("=" * 78)
    print(
        "RVO-ME AUDIT 02 - "
        "INTEGRITY VERIFICATION"
    )
    print("=" * 78)

    demo = summary[
        "demographics"
    ]

    print("\nPATIENT / EYE DEMOGRAPHICS")
    print("-" * 78)

    print(
        f"Unique patients:           "
        f"{demo['unique_patients']}"
    )

    print(
        f"Unique eyes:               "
        f"{demo['unique_eye_ids']}"
    )

    print(
        f"Patient-level gender:      "
        f"{demo['patient_gender_counts']}"
    )

    print(
        f"Demographic inconsistencies: "
        f"{demo['patient_demographic_inconsistencies']}"
    )

    split_summary = summary[
        "official_split_observed"
    ]

    print("\nOFFICIAL SPLIT - ACTUAL RELEASE")
    print("-" * 78)

    print(
        f"Train images:              "
        f"{split_summary['train_images']}"
    )

    print(
        f"Test images:               "
        f"{split_summary['test_images']}"
    )

    print(
        f"Train eyes:                "
        f"{split_summary['train_unique_eyes']}"
    )

    print(
        f"Test eyes:                 "
        f"{split_summary['test_unique_eyes']}"
    )

    print(
        f"Eye overlap:               "
        f"{split_summary['eye_overlap_count']}"
    )

    print(
        f"Overlapping Eye IDs:       "
        f"{split_summary['overlapping_eye_ids']}"
    )

    print(
        f"Train patients:            "
        f"{split_summary['train_unique_patients']}"
    )

    print(
        f"Test patients:             "
        f"{split_summary['test_unique_patients']}"
    )

    print(
        f"Patient overlap:           "
        f"{split_summary['patient_overlap_count']}"
    )

    print(
        f"Overlapping Patient IDs:   "
        f"{split_summary['overlapping_patient_ids']}"
    )

    print("\nPAPER-REPORTED EXPECTATION")
    print("-" * 78)

    print(
        "Training patients:         104"
    )

    print(
        "Test patients:             26"
    )

    print(
        "Expected patient overlap:  0"
    )

    if (
        split_summary[
            "patient_overlap_count"
        ] > 0
    ):

        print(
            "\n*** IMPORTANT: "
            "The downloaded split contains "
            "patient overlap. ***"
        )

    if (
        split_summary[
            "eye_overlap_count"
        ] > 0
    ):

        print(
            "*** IMPORTANT: "
            "At least one Eye ID occurs in "
            "both train and test. ***"
        )

    print("\nANNOTATION INTEGRITY")
    print("-" * 78)

    print(
        f"Total annotations:         "
        f"{summary['annotations']['total_annotations']}"
    )

    print(
        f"Status counts:             "
        f"{summary['annotations']['status_counts']}"
    )

    # Show only truly strange annotations among primary targets.
    primary_unusual = [
        row
        for row in
        annotations[
            "unusual_rows"
        ]
        if (
            row["status"]
            == "UNUSUAL"
        )
    ]

    print(
        f"Unusual primary annotations: "
        f"{len(primary_unusual)}"
    )

    if primary_unusual:

        reason_counts = Counter(
            row["reason"]
            for row in
            primary_unusual
        )

        print(
            "Unusual annotation reasons:"
        )

        for reason, count in (
            reason_counts.most_common()
        ):

            print(
                f"  {reason:<45} "
                f"{count}"
            )

    masks = summary[
        "mask_reconstruction"
    ]

    print("\nMASK RECONSTRUCTION")
    print("-" * 78)

    print(
        f"Masks compared:            "
        f"{masks['masks_compared']}"
    )

    print(
        f"Repository-exact matches:  "
        f"{masks['repo_exact_match_count']} "
        f"({masks['repo_exact_match_percent']:.4f}%)"
    )

    print(
        f"Repository-exact mismatch: "
        f"{masks['repo_exact_mismatch_count']}"
    )

    print(
        f"Target-only exact matches: "
        f"{masks['target_only_exact_match_count']} "
        f"({masks['target_only_match_percent']:.4f}%)"
    )

    print(
        f"Target-only mismatches:    "
        f"{masks['target_only_mismatch_count']}"
    )

    print(
        f"Missing mask inputs:       "
        f"{masks['missing_inputs']}"
    )

    print(
        f"Unexpected mask values:    "
        f"{masks['unexpected_official_mask_values']}"
    )

    print("\n")
    print("=" * 78)
    print("AUDIT 02 COMPLETE")
    print("=" * 78)

    print(
        "\nResults saved to:\n"
        f"{OUTPUT_DIR}"
    )

    print(
        "\nMost important outputs:"
    )

    print(
        "  1. audit02_summary.json"
    )

    print(
        "  2. split_patient_overlap.csv"
    )

    print(
        "  3. split_eye_overlap.csv"
    )

    print(
        "  4. unusual_annotations.csv"
    )

    print(
        "  5. mask_reconstruction_comparison.csv"
    )

    print(
        "  6. mask_class_agreement_repo_exact.csv"
    )

    print(
        "  7. mask_class_agreement_target_only.csv"
    )

    if SAVE_TOP_MISMATCH_IMAGES:

        print(
            "  8. top_mask_differences/"
        )


# =====================================================================
# MAIN
# =====================================================================

def main():

    print("=" * 78)
    print(
        "RVO-ME AUDIT 02 - "
        "DATA INTEGRITY VERIFICATION"
    )
    print("=" * 78)

    print(
        "\nDataset root:\n"
        f"{DATASET_ROOT}"
    )

    required_paths = [
        DATASET_ROOT,
        RAW_LABELME_DIR,
        IMAGE_SEG_DIR,
        SEG_MASKS_DIR,
        TRAIN_TXT,
        TEST_TXT,
    ]

    for path in required_paths:

        if not path.exists():

            print(
                "\nERROR: Required path "
                "not found:"
            )

            print(path)

            sys.exit(1)

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    # -------------------------------------------------------------
    # Build source lookup
    # -------------------------------------------------------------

    json_lookup, image_lookup = (
        build_raw_file_lookup()
    )

    print(
        f"Raw JSON files found:     "
        f"{len(json_lookup)}"
    )

    print(
        f"Raw images found:         "
        f"{len(image_lookup)}"
    )

    # -------------------------------------------------------------
    # Step 1
    # -------------------------------------------------------------

    demographics = (
        load_demographics()
    )

    patient_demo = (
        audit_patient_demographics(
            demographics
        )
    )

    # -------------------------------------------------------------
    # Step 2
    # -------------------------------------------------------------

    split = (
        audit_split_integrity(
            demographics
        )
    )

    # -------------------------------------------------------------
    # Step 3
    # -------------------------------------------------------------

    annotations = (
        audit_annotations(
            json_lookup
        )
    )

    # -------------------------------------------------------------
    # Step 4
    # -------------------------------------------------------------

    mask_audit = (
        audit_mask_reconstruction(
            json_lookup,
            image_lookup,
        )
    )

    # -------------------------------------------------------------
    # Reports
    # -------------------------------------------------------------

    summary = save_all_reports(
        demographics,
        patient_demo,
        split,
        annotations,
        mask_audit,
    )

    print_summary(
        summary,
        split,
        annotations,
    )


if __name__ == "__main__":
    main()