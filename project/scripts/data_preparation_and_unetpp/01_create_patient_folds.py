from __future__ import annotations

import csv
import hashlib
import json
import math
import re
import shutil
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import numpy as np
from openpyxl import load_workbook

try:
    from iterstrat.ml_stratifiers import (
        MultilabelStratifiedKFold,
        MultilabelStratifiedShuffleSplit,
    )
except ImportError:
    raise ImportError(
        "\nThe package 'iterative-stratification' is required.\n"
        "Install it using:\n\n"
        "    pip install iterative-stratification\n"
    )


# =====================================================================
# CONFIGURATION
# =====================================================================

DATASET_ROOT = Path(r"E:\Papers\RVO_ME\RVO-Lesion")

RAW_LABELME_DIR = DATASET_ROOT / "RVO_Lesion_Labelme"

IMAGE_SEG_DIR = DATASET_ROOT / "Image_Seg"
OFFICIAL_TRAIN_TXT = IMAGE_SEG_DIR / "train.txt"
OFFICIAL_TEST_TXT = IMAGE_SEG_DIR / "test.txt"

PROJECT_ROOT = DATASET_ROOT.parent / "project"

OUTPUT_DIR = (
    PROJECT_ROOT
    / "splits"
    / "patient_5fold_v1"
)

PRIMARY_LABELS = [
    "IRF",
    "SRF",
    "ELM",
    "EZ",
    "HF",
]

N_SPLITS = 5

# We search several valid stratified solutions and keep
# the one with the best overall balance.
N_CANDIDATE_SEEDS = 250

BASE_RANDOM_SEED = 20260814

# Each outer fold:
#
# ~26 patients test
# ~104 patients remaining
#
# From the remaining patients we reserve ~13 for validation.
VAL_PATIENTS_PER_FOLD = 13

# Protect the frozen split from accidental replacement.
FORCE_OVERWRITE = False

# Expected counts for the dataset version we audited.
STRICT_DATASET_CHECKS = True

EXPECTED_IMAGES = 3012
EXPECTED_PATIENTS = 130
EXPECTED_EYES = 146


# =====================================================================
# HELPERS
# =====================================================================

def safe_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def canonical_id(value: Any) -> str:

    if value is None:
        return ""

    if isinstance(value, int):
        return str(value)

    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)

    text = str(value).strip()

    try:
        number = float(text)

        if (
            number.is_integer()
            and "." in text
        ):
            return str(int(number))

    except ValueError:
        pass

    return text


def normalize_column_name(value: Any) -> str:

    return re.sub(
        r"[^a-z0-9]",
        "",
        safe_text(value).lower(),
    )


def natural_key(text: str):

    return [
        int(part)
        if part.isdigit()
        else part.lower()
        for part in re.split(
            r"(\d+)",
            str(text),
        )
    ]


def get_eye_id(filename: str) -> Optional[str]:

    stem = Path(filename).stem

    match = re.match(
        r"^([^_]+)_",
        stem,
    )

    if match:
        return canonical_id(
            match.group(1)
        )

    return None


def write_csv(
    path: Path,
    rows: List[Dict[str, Any]],
    fieldnames: Optional[List[str]] = None,
):

    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    if not rows:
        path.write_text(
            "",
            encoding="utf-8",
        )
        return

    if fieldnames is None:
        fieldnames = list(
            rows[0].keys()
        )

    with path.open(
        "w",
        newline="",
        encoding="utf-8-sig",
    ) as f:

        writer = csv.DictWriter(
            f,
            fieldnames=fieldnames,
        )

        writer.writeheader()
        writer.writerows(rows)


def write_lines(
    path: Path,
    values: List[str],
):

    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with path.open(
        "w",
        encoding="utf-8",
    ) as f:

        for value in values:
            f.write(
                str(value) + "\n"
            )


def save_json(
    path: Path,
    data: Any,
):

    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with path.open(
        "w",
        encoding="utf-8",
    ) as f:

        json.dump(
            data,
            f,
            indent=2,
            ensure_ascii=False,
            default=str,
        )


def sha256_file(path: Path) -> str:

    digest = hashlib.sha256()

    with path.open(
        "rb"
    ) as f:

        while True:

            block = f.read(
                1024 * 1024
            )

            if not block:
                break

            digest.update(block)

    return digest.hexdigest()


def read_name_list(
    path: Path,
) -> List[str]:

    if not path.exists():
        return []

    values = []

    with path.open(
        "r",
        encoding="utf-8-sig",
    ) as f:

        for line in f:

            line = line.strip()

            if line:
                values.append(
                    Path(line).name
                )

    return values


# =====================================================================
# FIND DEMOGRAPHICS FILE
# =====================================================================

def find_demographics_file() -> Path:

    candidates = list(
        DATASET_ROOT.glob("*.xlsx")
    )

    if not candidates:

        raise FileNotFoundError(
            "No XLSX file found in dataset root."
        )

    for path in candidates:

        if (
            "demograph"
            in path.name.lower()
        ):

            return path

    return candidates[0]


# =====================================================================
# LOAD DEMOGRAPHICS
# =====================================================================

def load_demographics():

    print(
        "\n[1/8] Loading patient demographics..."
    )

    path = find_demographics_file()

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    sheet = workbook.active

    rows = list(
        sheet.iter_rows(
            values_only=True
        )
    )

    if not rows:

        raise RuntimeError(
            "Demographics file is empty."
        )

    headers = [
        safe_text(v)
        for v in rows[0]
    ]

    normalized = [
        normalize_column_name(v)
        for v in headers
    ]

    columns = {}

    for idx, name in enumerate(
        normalized
    ):

        if name in {
            "patientid",
            "patient",
        }:

            columns[
                "patient"
            ] = idx

        elif name in {
            "eyeid",
            "eye",
        }:

            columns[
                "eye"
            ] = idx

        elif name.startswith(
            "age"
        ):

            columns[
                "age"
            ] = idx

        elif name in {
            "gender",
            "sex",
        }:

            columns[
                "gender"
            ] = idx

        elif (
            "eyecategory"
            in name
            or "eyecateg"
            in name
            or name
            == "laterality"
        ):

            columns[
                "eye_category"
            ] = idx

    for required in [
        "patient",
        "eye",
        "age",
        "gender",
    ]:

        if required not in columns:

            raise RuntimeError(
                "Could not identify "
                f"{required} column."
            )

    eye_to_patient = {}

    patient_info = defaultdict(
        lambda: {
            "eyes": set(),
            "ages": set(),
            "genders": set(),
            "eye_categories": set(),
        }
    )

    for values in rows[1:]:

        if not any(
            value is not None
            for value in values
        ):
            continue

        patient_id = canonical_id(
            values[
                columns["patient"]
            ]
        )

        eye_id = canonical_id(
            values[
                columns["eye"]
            ]
        )

        age_value = values[
            columns["age"]
        ]

        gender = safe_text(
            values[
                columns["gender"]
            ]
        )

        category = ""

        if (
            "eye_category"
            in columns
        ):

            category = safe_text(
                values[
                    columns[
                        "eye_category"
                    ]
                ]
            )

        if not patient_id:
            continue

        if not eye_id:
            continue

        if (
            eye_id
            in eye_to_patient
            and eye_to_patient[
                eye_id
            ]
            != patient_id
        ):

            raise RuntimeError(
                f"Eye {eye_id} maps "
                "to multiple patients."
            )

        eye_to_patient[
            eye_id
        ] = patient_id

        patient_info[
            patient_id
        ]["eyes"].add(
            eye_id
        )

        if age_value is not None:

            patient_info[
                patient_id
            ]["ages"].add(
                float(age_value)
            )

        if gender:

            patient_info[
                patient_id
            ]["genders"].add(
                gender
            )

        if category:

            patient_info[
                patient_id
            ][
                "eye_categories"
            ].add(
                category
            )

    workbook.close()

    clean_patients = {}

    for patient_id, info in (
        patient_info.items()
    ):

        if len(
            info["ages"]
        ) != 1:

            raise RuntimeError(
                f"Patient {patient_id} "
                "has inconsistent age."
            )

        if len(
            info["genders"]
        ) != 1:

            raise RuntimeError(
                f"Patient {patient_id} "
                "has inconsistent gender."
            )

        clean_patients[
            patient_id
        ] = {
            "patient_id":
                patient_id,

            "eyes":
                sorted(
                    info["eyes"],
                    key=natural_key,
                ),

            "age":
                list(
                    info["ages"]
                )[0],

            "gender":
                list(
                    info["genders"]
                )[0],

            "eye_categories":
                sorted(
                    info[
                        "eye_categories"
                    ]
                ),
        }

    return (
        clean_patients,
        eye_to_patient,
        path,
    )


# =====================================================================
# BUILD IMAGE MANIFEST
# =====================================================================

def build_image_manifest(
    eye_to_patient
):

    print(
        "[2/8] Building raw image / annotation manifest..."
    )

    json_files = sorted(
        RAW_LABELME_DIR.rglob(
            "*.json"
        ),
        key=lambda p:
            natural_key(p.stem),
    )

    image_rows = []

    seen_stems = set()

    for index, json_path in enumerate(
        json_files,
        start=1,
    ):

        stem = json_path.stem

        if stem in seen_stems:

            raise RuntimeError(
                f"Duplicate image stem: {stem}"
            )

        seen_stems.add(
            stem
        )

        eye_id = get_eye_id(
            stem
        )

        if not eye_id:

            raise RuntimeError(
                "Cannot extract Eye ID "
                f"from {stem}"
            )

        patient_id = (
            eye_to_patient.get(
                eye_id
            )
        )

        if not patient_id:

            raise RuntimeError(
                "Eye ID not present in "
                f"demographics: {eye_id}"
            )

        image_path = None

        for ext in [
            ".jpg",
            ".jpeg",
            ".png",
            ".bmp",
            ".tif",
            ".tiff",
        ]:

            candidate = (
                json_path.with_suffix(
                    ext
                )
            )

            if candidate.exists():

                image_path = (
                    candidate
                )

                break

        if image_path is None:

            raise RuntimeError(
                "Raw image missing for "
                f"{json_path}"
            )

        with json_path.open(
            "r",
            encoding="utf-8-sig",
        ) as f:

            data = json.load(f)

        label_counts = Counter()

        labels_in_image = set()

        for shape in data.get(
            "shapes",
            []
        ):

            label = safe_text(
                shape.get(
                    "label"
                )
            )

            if (
                label
                in PRIMARY_LABELS
            ):

                label_counts[
                    label
                ] += 1

                labels_in_image.add(
                    label
                )

        row = {
            "image_filename":
                image_path.name,

            "image_stem":
                stem,

            "patient_id":
                patient_id,

            "eye_id":
                eye_id,

            "raw_image_path":
                str(image_path),

            "json_path":
                str(json_path),

            "labels":
                ";".join(
                    sorted(
                        labels_in_image
                    )
                ),
        }

        for label in PRIMARY_LABELS:

            row[
                f"{label}_present"
            ] = int(
                label
                in labels_in_image
            )

            row[
                f"{label}_annotations"
            ] = int(
                label_counts[
                    label
                ]
            )

        image_rows.append(
            row
        )

        if index % 500 == 0:

            print(
                f"    Processed "
                f"{index}/"
                f"{len(json_files)}"
            )

    return image_rows


# =====================================================================
# BUILD PATIENT MANIFEST
# =====================================================================

def build_patient_manifest(
    demographics,
    image_rows,
):

    print(
        "[3/8] Aggregating patient-level biomarker statistics..."
    )

    rows_by_patient = defaultdict(
        list
    )

    for row in image_rows:

        rows_by_patient[
            row[
                "patient_id"
            ]
        ].append(
            row
        )

    patient_rows = []

    for patient_id in sorted(
        demographics,
        key=natural_key,
    ):

        demo = demographics[
            patient_id
        ]

        images = rows_by_patient[
            patient_id
        ]

        if not images:

            raise RuntimeError(
                f"Patient {patient_id} "
                "has no images."
            )

        row = {
            "patient_id":
                patient_id,

            "age":
                float(
                    demo["age"]
                ),

            "gender":
                demo["gender"],

            "n_eyes":
                len(
                    demo["eyes"]
                ),

            "eye_ids":
                ";".join(
                    demo["eyes"]
                ),

            "bilateral":
                int(
                    len(
                        demo["eyes"]
                    ) > 1
                ),

            "n_images":
                len(images),
        }

        for label in PRIMARY_LABELS:

            positive_images = sum(
                int(
                    image[
                        f"{label}_present"
                    ]
                )
                for image in images
            )

            annotation_count = sum(
                int(
                    image[
                        f"{label}_annotations"
                    ]
                )
                for image in images
            )

            row[
                f"{label}_present"
            ] = int(
                positive_images > 0
            )

            row[
                f"{label}_positive_images"
            ] = positive_images

            row[
                f"{label}_positive_rate"
            ] = (
                positive_images
                / len(images)
            )

            row[
                f"{label}_annotations"
            ] = annotation_count

        patient_rows.append(
            row
        )

    return patient_rows


# =====================================================================
# DATASET INTEGRITY CHECK
# =====================================================================

def verify_dataset(
    patient_rows,
    image_rows,
    eye_to_patient,
):

    print(
        "[4/8] Verifying dataset counts..."
    )

    n_patients = len(
        patient_rows
    )

    n_images = len(
        image_rows
    )

    n_eyes = len(
        eye_to_patient
    )

    print(
        f"    Patients: {n_patients}"
    )

    print(
        f"    Eyes:     {n_eyes}"
    )

    print(
        f"    Images:   {n_images}"
    )

    if STRICT_DATASET_CHECKS:

        if (
            n_patients
            != EXPECTED_PATIENTS
        ):

            raise RuntimeError(
                "Expected "
                f"{EXPECTED_PATIENTS} "
                f"patients, found "
                f"{n_patients}."
            )

        if (
            n_eyes
            != EXPECTED_EYES
        ):

            raise RuntimeError(
                "Expected "
                f"{EXPECTED_EYES} eyes, "
                f"found {n_eyes}."
            )

        if (
            n_images
            != EXPECTED_IMAGES
        ):

            raise RuntimeError(
                "Expected "
                f"{EXPECTED_IMAGES} images, "
                f"found {n_images}."
            )


# =====================================================================
# CREATE STRATIFICATION MATRIX
# =====================================================================

def create_stratification_matrix(
    patient_rows
):

    print(
        "[5/8] Creating multilabel stratification features..."
    )

    feature_columns = []

    feature_vectors = []

    # -------------------------------------------------------------
    # Calculate burden thresholds
    # -------------------------------------------------------------

    burden_thresholds = {}

    for label in PRIMARY_LABELS:

        positive_rates = [
            float(
                row[
                    f"{label}_positive_rate"
                ]
            )
            for row in patient_rows
            if int(
                row[
                    f"{label}_present"
                ]
            ) == 1
        ]

        if positive_rates:

            burden_thresholds[
                label
            ] = float(
                np.median(
                    positive_rates
                )
            )

        else:

            burden_thresholds[
                label
            ] = 0.0

    # -------------------------------------------------------------
    # Age quartiles
    # -------------------------------------------------------------

    ages = np.array(
        [
            float(
                row["age"]
            )
            for row
            in patient_rows
        ],
        dtype=float,
    )

    age_thresholds = np.quantile(
        ages,
        [
            0.25,
            0.50,
            0.75,
        ],
    )

    # -------------------------------------------------------------
    # Feature names
    # -------------------------------------------------------------

    for label in PRIMARY_LABELS:

        feature_columns.append(
            f"{label}_present"
        )

    for label in PRIMARY_LABELS:

        feature_columns.append(
            f"{label}_high_burden"
        )

    feature_columns.extend([
        "female",
        "bilateral",
        "age_q1",
        "age_q2",
        "age_q3",
        "age_q4",
    ])

    # -------------------------------------------------------------
    # Build feature matrix
    # -------------------------------------------------------------

    for row in patient_rows:

        features = []

        for label in PRIMARY_LABELS:

            features.append(
                int(
                    row[
                        f"{label}_present"
                    ]
                )
            )

        for label in PRIMARY_LABELS:

            rate = float(
                row[
                    f"{label}_positive_rate"
                ]
            )

            present = int(
                row[
                    f"{label}_present"
                ]
            )

            high = int(
                present == 1
                and rate
                >= burden_thresholds[
                    label
                ]
            )

            features.append(
                high
            )

        gender = (
            safe_text(
                row["gender"]
            ).lower()
        )

        features.append(
            int(
                gender.startswith(
                    "f"
                )
            )
        )

        features.append(
            int(
                row[
                    "bilateral"
                ]
            )
        )

        age = float(
            row["age"]
        )

        age_bin = int(
            np.digitize(
                age,
                age_thresholds,
                right=True,
            )
        )

        for q in range(4):

            features.append(
                int(
                    age_bin == q
                )
            )

        feature_vectors.append(
            features
        )

    y = np.asarray(
        feature_vectors,
        dtype=np.int32,
    )

    return (
        y,
        feature_columns,
        burden_thresholds,
        age_thresholds,
    )


# =====================================================================
# BALANCE SCORING
# =====================================================================

def calculate_fold_score(
    fold_test_indices,
    patient_rows,
):

    """
    Lower score = better balance.

    This is NOT model-performance optimization.

    It only chooses among valid stratified partitions based on
    patient/image/biomarker/demographic balance.
    """

    n_folds = len(
        fold_test_indices
    )

    score_sum = 0.0
    weight_sum = 0.0

    def score_additive(
        values,
        weight,
    ):

        nonlocal score_sum
        nonlocal weight_sum

        values = np.asarray(
            values,
            dtype=float,
        )

        target = (
            values.sum()
            / n_folds
        )

        denominator = max(
            abs(target),
            1.0,
        )

        deviations = []

        for indices in (
            fold_test_indices
        ):

            observed = values[
                indices
            ].sum()

            deviations.append(
                (
                    observed
                    - target
                )
                / denominator
            )

        value_score = float(
            np.mean(
                np.square(
                    deviations
                )
            )
        )

        score_sum += (
            weight
            * value_score
        )

        weight_sum += weight

    # ---------------------------------------------------------
    # Basic dataset balance
    # ---------------------------------------------------------

    score_additive(
        [
            row["n_eyes"]
            for row
            in patient_rows
        ],
        weight=1.0,
    )

    score_additive(
        [
            row["n_images"]
            for row
            in patient_rows
        ],
        weight=1.5,
    )

    score_additive(
        [
            1
            if safe_text(
                row["gender"]
            ).lower().startswith(
                "f"
            )
            else 0
            for row
            in patient_rows
        ],
        weight=1.0,
    )

    score_additive(
        [
            row["bilateral"]
            for row
            in patient_rows
        ],
        weight=1.0,
    )

    # ---------------------------------------------------------
    # Biomarker balance
    # ---------------------------------------------------------

    for label in PRIMARY_LABELS:

        score_additive(
            [
                row[
                    f"{label}_present"
                ]
                for row
                in patient_rows
            ],
            weight=3.0,
        )

        score_additive(
            [
                row[
                    f"{label}_positive_images"
                ]
                for row
                in patient_rows
            ],
            weight=2.0,
        )

        # Log transform annotation counts so a single patient
        # with many IRF components cannot dominate the score.
        score_additive(
            [
                math.log1p(
                    row[
                        f"{label}_annotations"
                    ]
                )
                for row
                in patient_rows
            ],
            weight=1.0,
        )

    # ---------------------------------------------------------
    # Age mean balance
    # ---------------------------------------------------------

    ages = np.asarray(
        [
            row["age"]
            for row
            in patient_rows
        ],
        dtype=float,
    )

    global_age_mean = float(
        ages.mean()
    )

    global_age_std = max(
        float(
            ages.std()
        ),
        1.0,
    )

    age_deviation = []

    for indices in (
        fold_test_indices
    ):

        fold_mean = float(
            ages[
                indices
            ].mean()
        )

        age_deviation.append(
            (
                fold_mean
                - global_age_mean
            )
            / global_age_std
        )

    score_sum += float(
        np.mean(
            np.square(
                age_deviation
            )
        )
    )

    weight_sum += 1.0

    return (
        score_sum
        / weight_sum
    )


# =====================================================================
# SEARCH FOR BEST 5-FOLD ASSIGNMENT
# =====================================================================

def create_outer_folds(
    patient_rows,
    y,
):

    print(
        "[6/8] Searching for balanced patient-disjoint 5-fold split..."
    )

    x_dummy = np.zeros(
        (
            len(patient_rows),
            1,
        ),
        dtype=np.float32,
    )

    best_score = None
    best_seed = None
    best_test_indices = None

    for candidate in range(
        N_CANDIDATE_SEEDS
    ):

        seed = (
            BASE_RANDOM_SEED
            + candidate
        )

        splitter = (
            MultilabelStratifiedKFold(
                n_splits=N_SPLITS,
                shuffle=True,
                random_state=seed,
            )
        )

        test_indices = []

        for _, test_idx in (
            splitter.split(
                x_dummy,
                y,
            )
        ):

            test_indices.append(
                np.asarray(
                    test_idx,
                    dtype=int,
                )
            )

        score = (
            calculate_fold_score(
                test_indices,
                patient_rows,
            )
        )

        if (
            best_score is None
            or score
            < best_score
        ):

            best_score = score
            best_seed = seed
            best_test_indices = (
                test_indices
            )

    print(
        "    Selected seed: "
        f"{best_seed}"
    )

    print(
        "    Balance score: "
        f"{best_score:.8f}"
    )

    return (
        best_test_indices,
        best_seed,
        best_score,
    )


# =====================================================================
# CREATE VALIDATION SUBSET
# =====================================================================

def create_validation_set(
    train_pool_indices,
    y,
    seed,
):

    pool = np.asarray(
        train_pool_indices,
        dtype=int,
    )

    y_pool = y[
        pool
    ]

    x_dummy = np.zeros(
        (
            len(pool),
            1,
        ),
        dtype=np.float32,
    )

    splitter = (
        MultilabelStratifiedShuffleSplit(
            n_splits=1,
            test_size=
                VAL_PATIENTS_PER_FOLD,
            random_state=seed,
        )
    )

    relative_train, relative_val = next(
        splitter.split(
            x_dummy,
            y_pool,
        )
    )

    train_indices = pool[
        relative_train
    ]

    val_indices = pool[
        relative_val
    ]

    return (
        train_indices,
        val_indices,
    )


# =====================================================================
# SPLIT SUMMARY
# =====================================================================

def summarize_patients(
    indices,
    patient_rows,
):

    selected = [
        patient_rows[
            int(i)
        ]
        for i in indices
    ]

    result = {
        "patients":
            len(selected),

        "eyes":
            sum(
                row["n_eyes"]
                for row
                in selected
            ),

        "images":
            sum(
                row["n_images"]
                for row
                in selected
            ),

        "mean_age":
            float(
                np.mean(
                    [
                        row["age"]
                        for row
                        in selected
                    ]
                )
            ),

        "female":
            sum(
                1
                for row
                in selected
                if safe_text(
                    row["gender"]
                ).lower().startswith(
                    "f"
                )
            ),
    }

    result[
        "male"
    ] = (
        result["patients"]
        - result["female"]
    )

    for label in PRIMARY_LABELS:

        result[
            f"{label}_positive_patients"
        ] = sum(
            row[
                f"{label}_present"
            ]
            for row
            in selected
        )

        result[
            f"{label}_positive_images"
        ] = sum(
            row[
                f"{label}_positive_images"
            ]
            for row
            in selected
        )

        result[
            f"{label}_annotations"
        ] = sum(
            row[
                f"{label}_annotations"
            ]
            for row
            in selected
        )

    return result


# =====================================================================
# VERIFY ZERO LEAKAGE
# =====================================================================

def verify_role_disjointness(
    train_indices,
    val_indices,
    test_indices,
    patient_rows,
):

    def get_patients(
        indices
    ):

        return {
            patient_rows[
                int(i)
            ]["patient_id"]
            for i in indices
        }

    def get_eyes(
        indices
    ):

        eyes = set()

        for i in indices:

            eyes.update(
                patient_rows[
                    int(i)
                ]["eye_ids"].split(
                    ";"
                )
            )

        return {
            eye
            for eye in eyes
            if eye
        }

    train_patients = (
        get_patients(
            train_indices
        )
    )

    val_patients = (
        get_patients(
            val_indices
        )
    )

    test_patients = (
        get_patients(
            test_indices
        )
    )

    if (
        train_patients
        & val_patients
    ):

        raise RuntimeError(
            "Patient leakage: train/val."
        )

    if (
        train_patients
        & test_patients
    ):

        raise RuntimeError(
            "Patient leakage: train/test."
        )

    if (
        val_patients
        & test_patients
    ):

        raise RuntimeError(
            "Patient leakage: val/test."
        )

    train_eyes = (
        get_eyes(
            train_indices
        )
    )

    val_eyes = (
        get_eyes(
            val_indices
        )
    )

    test_eyes = (
        get_eyes(
            test_indices
        )
    )

    if train_eyes & val_eyes:

        raise RuntimeError(
            "Eye leakage: train/val."
        )

    if train_eyes & test_eyes:

        raise RuntimeError(
            "Eye leakage: train/test."
        )

    if val_eyes & test_eyes:

        raise RuntimeError(
            "Eye leakage: val/test."
        )

    return {
        "train_patients":
            len(train_patients),

        "val_patients":
            len(val_patients),

        "test_patients":
            len(test_patients),

        "train_eyes":
            len(train_eyes),

        "val_eyes":
            len(val_eyes),

        "test_eyes":
            len(test_eyes),

        "patient_leakage":
            0,

        "eye_leakage":
            0,
    }


# =====================================================================
# SAVE SPLITS
# =====================================================================

def save_splits(
    patient_rows,
    image_rows,
    y,
    feature_columns,
    burden_thresholds,
    age_thresholds,
    outer_test_indices,
    selected_seed,
    selected_score,
    demographics_path,
):

    print(
        "[7/8] Saving frozen research splits..."
    )

    if OUTPUT_DIR.exists():

        if (
            any(
                OUTPUT_DIR.iterdir()
            )
            and not FORCE_OVERWRITE
        ):

            raise RuntimeError(
                "\nOutput directory already "
                "contains files:\n\n"
                f"{OUTPUT_DIR}\n\n"
                "This protection prevents accidental "
                "replacement of the frozen research split.\n"
                "If you intentionally want to regenerate it, "
                "either delete the directory or set:\n\n"
                "FORCE_OVERWRITE = True\n"
            )

        if FORCE_OVERWRITE:

            shutil.rmtree(
                OUTPUT_DIR
            )

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    # ---------------------------------------------------------
    # Official split lookup, only for reference.
    # ---------------------------------------------------------

    official_train = set(
        read_name_list(
            OFFICIAL_TRAIN_TXT
        )
    )

    official_test = set(
        read_name_list(
            OFFICIAL_TEST_TXT
        )
    )

    for row in image_rows:

        filename = row[
            "image_filename"
        ]

        if filename in (
            official_train
        ):

            row[
                "official_split"
            ] = "train"

        elif filename in (
            official_test
        ):

            row[
                "official_split"
            ] = "test"

        else:

            row[
                "official_split"
            ] = "unknown"

    # ---------------------------------------------------------
    # Outer-fold patient assignment
    # ---------------------------------------------------------

    patient_test_fold = {}

    for fold_idx, indices in enumerate(
        outer_test_indices,
        start=1,
    ):

        for index in indices:

            patient_id = (
                patient_rows[
                    int(index)
                ][
                    "patient_id"
                ]
            )

            if (
                patient_id
                in patient_test_fold
            ):

                raise RuntimeError(
                    "A patient was assigned "
                    "to more than one test fold."
                )

            patient_test_fold[
                patient_id
            ] = fold_idx

    if (
        len(
            patient_test_fold
        )
        != len(
            patient_rows
        )
    ):

        raise RuntimeError(
            "Not every patient received "
            "exactly one outer test fold."
        )

    # ---------------------------------------------------------
    # Save patient manifest
    # ---------------------------------------------------------

    patient_manifest = []

    for row in patient_rows:

        output_row = dict(
            row
        )

        output_row[
            "outer_test_fold"
        ] = patient_test_fold[
            row["patient_id"]
        ]

        patient_manifest.append(
            output_row
        )

    patient_manifest_path = (
        OUTPUT_DIR
        / "patient_manifest.csv"
    )

    write_csv(
        patient_manifest_path,
        patient_manifest,
    )

    # ---------------------------------------------------------
    # Save image manifest
    # ---------------------------------------------------------

    image_manifest = []

    for row in image_rows:

        output_row = dict(
            row
        )

        output_row[
            "outer_test_fold"
        ] = patient_test_fold[
            row["patient_id"]
        ]

        image_manifest.append(
            output_row
        )

    write_csv(
        OUTPUT_DIR
        / "image_manifest.csv",
        image_manifest,
    )

    # ---------------------------------------------------------
    # Save outer assignment
    # ---------------------------------------------------------

    fold_assignment_rows = []

    for row in patient_rows:

        fold_assignment_rows.append({
            "patient_id":
                row[
                    "patient_id"
                ],

            "eye_ids":
                row[
                    "eye_ids"
                ],

            "outer_test_fold":
                patient_test_fold[
                    row[
                        "patient_id"
                    ]
                ],
        })

    fold_assignment_path = (
        OUTPUT_DIR
        / "fold_assignments.csv"
    )

    write_csv(
        fold_assignment_path,
        fold_assignment_rows,
    )

    # ---------------------------------------------------------
    # Create train / validation / test for each fold
    # ---------------------------------------------------------

    role_balance_rows = []
    test_balance_rows = []

    fold_integrity = []

    all_indices = np.arange(
        len(
            patient_rows
        )
    )

    for fold_idx, test_indices in enumerate(
        outer_test_indices,
        start=1,
    ):

        test_indices = np.asarray(
            test_indices,
            dtype=int,
        )

        test_set = set(
            test_indices.tolist()
        )

        train_pool_indices = np.asarray(
            [
                i
                for i in all_indices
                if int(i)
                not in test_set
            ],
            dtype=int,
        )

        train_indices, val_indices = (
            create_validation_set(
                train_pool_indices,
                y,
                seed=(
                    selected_seed
                    + 10000
                    + fold_idx
                ),
            )
        )

        integrity = (
            verify_role_disjointness(
                train_indices,
                val_indices,
                test_indices,
                patient_rows,
            )
        )

        integrity[
            "fold"
        ] = fold_idx

        fold_integrity.append(
            integrity
        )

        fold_dir = (
            OUTPUT_DIR
            / f"fold_{fold_idx}"
        )

        fold_dir.mkdir(
            parents=True,
            exist_ok=True,
        )

        role_indices = {
            "train":
                train_indices,

            "val":
                val_indices,

            "test":
                test_indices,
        }

        patient_role_rows = []

        patient_to_role = {}

        for role, indices in (
            role_indices.items()
        ):

            patient_ids = sorted(
                [
                    patient_rows[
                        int(i)
                    ]["patient_id"]
                    for i in indices
                ],
                key=natural_key,
            )

            write_lines(
                fold_dir
                / f"{role}_patients.txt",
                patient_ids,
            )

            for patient_id in (
                patient_ids
            ):

                patient_to_role[
                    patient_id
                ] = role

        for row in patient_rows:

            patient_role_rows.append({
                "patient_id":
                    row["patient_id"],

                "eye_ids":
                    row["eye_ids"],

                "role":
                    patient_to_role[
                        row[
                            "patient_id"
                        ]
                    ],
            })

        write_csv(
            fold_dir
            / "patient_roles.csv",
            patient_role_rows,
        )

        # -----------------------------------------------------
        # Image files
        # -----------------------------------------------------

        role_images = {
            "train": [],
            "val": [],
            "test": [],
        }

        for image in image_rows:

            role = (
                patient_to_role[
                    image[
                        "patient_id"
                    ]
                ]
            )

            role_images[
                role
            ].append(
                image[
                    "image_filename"
                ]
            )

        for role in [
            "train",
            "val",
            "test",
        ]:

            filenames = sorted(
                role_images[
                    role
                ],
                key=natural_key,
            )

            write_lines(
                fold_dir
                / f"{role}.txt",
                filenames,
            )

        # -----------------------------------------------------
        # Role balance
        # -----------------------------------------------------

        for role, indices in (
            role_indices.items()
        ):

            summary = (
                summarize_patients(
                    indices,
                    patient_rows,
                )
            )

            row = {
                "fold":
                    fold_idx,

                "role":
                    role,
            }

            row.update(
                summary
            )

            role_balance_rows.append(
                row
            )

        test_summary = (
            summarize_patients(
                test_indices,
                patient_rows,
            )
        )

        test_row = {
            "fold":
                fold_idx
        }

        test_row.update(
            test_summary
        )

        test_balance_rows.append(
            test_row
        )

    # ---------------------------------------------------------
    # Save balance / integrity reports
    # ---------------------------------------------------------

    write_csv(
        OUTPUT_DIR
        / "fold_role_balance.csv",
        role_balance_rows,
    )

    write_csv(
        OUTPUT_DIR
        / "outer_test_fold_balance.csv",
        test_balance_rows,
    )

    write_csv(
        OUTPUT_DIR
        / "fold_integrity.csv",
        fold_integrity,
    )

    # ---------------------------------------------------------
    # Stratification features
    # ---------------------------------------------------------

    stratification_rows = []

    for index, row in enumerate(
        patient_rows
    ):

        output_row = {
            "patient_id":
                row[
                    "patient_id"
                ]
        }

        for feature_idx, name in enumerate(
            feature_columns
        ):

            output_row[
                name
            ] = int(
                y[
                    index,
                    feature_idx,
                ]
            )

        stratification_rows.append(
            output_row
        )

    write_csv(
        OUTPUT_DIR
        / "stratification_features.csv",
        stratification_rows,
    )

    # ---------------------------------------------------------
    # Fingerprint
    # ---------------------------------------------------------

    assignment_sha256 = (
        sha256_file(
            fold_assignment_path
        )
    )

    # ---------------------------------------------------------
    # Summary JSON
    # ---------------------------------------------------------

    summary = {
        "split_name":
            "patient_5fold_v1",

        "dataset_root":
            str(
                DATASET_ROOT
            ),

        "demographics_file":
            str(
                demographics_path
            ),

        "n_patients":
            len(
                patient_rows
            ),

        "n_images":
            len(
                image_rows
            ),

        "n_outer_folds":
            N_SPLITS,

        "validation_patients_per_fold":
            VAL_PATIENTS_PER_FOLD,

        "candidate_seeds_tested":
            N_CANDIDATE_SEEDS,

        "selected_random_seed":
            selected_seed,

        "selected_balance_score":
            selected_score,

        "primary_labels":
            PRIMARY_LABELS,

        "stratification_features":
            feature_columns,

        "burden_thresholds":
            burden_thresholds,

        "age_quartile_thresholds":
            [
                float(x)
                for x
                in age_thresholds
            ],

        "fold_assignment_sha256":
            assignment_sha256,

        "fold_integrity":
            fold_integrity,

        "important_rule":
            (
                "All images and eyes belonging "
                "to the same patient remain "
                "within the same role for every "
                "train/validation/test split."
            ),
    }

    save_json(
        OUTPUT_DIR
        / "split_summary.json",
        summary,
    )

    return (
        summary,
        test_balance_rows,
        role_balance_rows,
    )


# =====================================================================
# FINAL VALIDATION
# =====================================================================

def final_validation(
    patient_rows,
    image_rows,
    outer_test_indices,
):

    print(
        "[8/8] Running final split integrity checks..."
    )

    patient_test_counter = Counter()

    image_test_counter = Counter()

    patient_images = defaultdict(
        list
    )

    for image in image_rows:

        patient_images[
            image[
                "patient_id"
            ]
        ].append(
            image[
                "image_filename"
            ]
        )

    for indices in (
        outer_test_indices
    ):

        for idx in indices:

            patient_id = (
                patient_rows[
                    int(idx)
                ][
                    "patient_id"
                ]
            )

            patient_test_counter[
                patient_id
            ] += 1

            for filename in (
                patient_images[
                    patient_id
                ]
            ):

                image_test_counter[
                    filename
                ] += 1

    bad_patients = {
        patient:
            count
        for patient, count
        in patient_test_counter.items()
        if count != 1
    }

    bad_images = {
        image:
            count
        for image, count
        in image_test_counter.items()
        if count != 1
    }

    if bad_patients:

        raise RuntimeError(
            "Some patients do not occur "
            "exactly once as outer test:"
            f"\n{bad_patients}"
        )

    if bad_images:

        raise RuntimeError(
            "Some images do not occur "
            "exactly once as outer test."
        )

    if (
        len(
            patient_test_counter
        )
        != len(
            patient_rows
        )
    ):

        raise RuntimeError(
            "Not every patient appeared "
            "in outer testing."
        )

    if (
        len(
            image_test_counter
        )
        != len(
            image_rows
        )
    ):

        raise RuntimeError(
            "Not every image appeared "
            "in outer testing."
        )

    print(
        "    Zero patient leakage: PASS"
    )

    print(
        "    Zero eye leakage:     PASS"
    )

    print(
        "    Every patient tested exactly once: PASS"
    )

    print(
        "    Every image tested exactly once:   PASS"
    )


# =====================================================================
# TERMINAL REPORT
# =====================================================================

def print_final_report(
    summary,
    test_balance_rows,
):

    print("\n")
    print("=" * 90)
    print(
        "RVO-ME PATIENT-DISJOINT "
        "5-FOLD SPLIT COMPLETE"
    )
    print("=" * 90)

    print(
        f"\nSelected seed: "
        f"{summary['selected_random_seed']}"
    )

    print(
        "Balance score: "
        f"{summary['selected_balance_score']:.8f}"
    )

    print(
        "\nOuter test-fold balance:"
    )

    print("-" * 90)

    header = (
        f"{'Fold':<6}"
        f"{'Patients':<10}"
        f"{'Eyes':<8}"
        f"{'Images':<9}"
        f"{'Female':<9}"
        f"{'AgeMean':<10}"
        f"{'IRF':<7}"
        f"{'SRF':<7}"
        f"{'ELM':<7}"
        f"{'EZ':<7}"
        f"{'HF':<7}"
    )

    print(header)

    print("-" * 90)

    for row in test_balance_rows:

        print(
            f"{row['fold']:<6}"
            f"{row['patients']:<10}"
            f"{row['eyes']:<8}"
            f"{row['images']:<9}"
            f"{row['female']:<9}"
            f"{row['mean_age']:<10.2f}"
            f"{row['IRF_positive_patients']:<7}"
            f"{row['SRF_positive_patients']:<7}"
            f"{row['ELM_positive_patients']:<7}"
            f"{row['EZ_positive_patients']:<7}"
            f"{row['HF_positive_patients']:<7}"
        )

    print("\nIntegrity:")

    for item in (
        summary[
            "fold_integrity"
        ]
    ):

        print(
            f"  Fold {item['fold']}: "
            f"train={item['train_patients']} patients, "
            f"val={item['val_patients']}, "
            f"test={item['test_patients']}, "
            f"patient leakage={item['patient_leakage']}, "
            f"eye leakage={item['eye_leakage']}"
        )

    print(
        "\nFold assignment SHA-256:"
    )

    print(
        summary[
            "fold_assignment_sha256"
        ]
    )

    print(
        "\nFrozen split saved to:"
    )

    print(
        OUTPUT_DIR
    )

    print(
        "\nIMPORTANT:"
    )

    print(
        "Do not regenerate this split after "
        "we start model experiments."
    )

    print(
        "The SHA-256 fingerprint above identifies "
        "the exact patient assignment used in the study."
    )

    print("\n")
    print("=" * 90)


# =====================================================================
# MAIN
# =====================================================================

def main():

    print("=" * 90)

    print(
        "RVO-ME: CREATE PATIENT-DISJOINT "
        "5-FOLD RESEARCH SPLIT"
    )

    print("=" * 90)

    print(
        f"\nDataset:\n{DATASET_ROOT}"
    )

    if not DATASET_ROOT.exists():

        raise FileNotFoundError(
            DATASET_ROOT
        )

    if not RAW_LABELME_DIR.exists():

        raise FileNotFoundError(
            RAW_LABELME_DIR
        )

    demographics, eye_to_patient, demographics_path = (
        load_demographics()
    )

    image_rows = (
        build_image_manifest(
            eye_to_patient
        )
    )

    patient_rows = (
        build_patient_manifest(
            demographics,
            image_rows,
        )
    )

    verify_dataset(
        patient_rows,
        image_rows,
        eye_to_patient,
    )

    (
        y,
        feature_columns,
        burden_thresholds,
        age_thresholds,
    ) = create_stratification_matrix(
        patient_rows
    )

    (
        outer_test_indices,
        selected_seed,
        selected_score,
    ) = create_outer_folds(
        patient_rows,
        y,
    )

    final_validation(
        patient_rows,
        image_rows,
        outer_test_indices,
    )

    (
        summary,
        test_balance_rows,
        role_balance_rows,
    ) = save_splits(
        patient_rows,
        image_rows,
        y,
        feature_columns,
        burden_thresholds,
        age_thresholds,
        outer_test_indices,
        selected_seed,
        selected_score,
        demographics_path,
    )

    print_final_report(
        summary,
        test_balance_rows,
    )


if __name__ == "__main__":
    main()