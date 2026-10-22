from __future__ import annotations

import csv
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from statistics import mean, median

# ============================================================
# CONFIGURATION
# ============================================================

DATASET_ROOT = Path(r"E:\Papers\RVO_ME\RVO-Lesion")

RAW_LABELME_DIR = DATASET_ROOT / "RVO_Lesion_Labelme"
IMAGE_SEG_DIR = DATASET_ROOT / "Image_Seg"
HF_DETECTION_DIR = DATASET_ROOT / "RVO_HF_Detection"

OUTPUT_DIR = DATASET_ROOT.parent / "audit_results"

# Do NOT modify the source dataset.
# All generated reports will be written to OUTPUT_DIR.


# ============================================================
# OPTIONAL DEPENDENCIES
# ============================================================

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

try:
    from PIL import Image
except ImportError:
    Image = None


# ============================================================
# HELPERS
# ============================================================

IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff"}


def safe_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_column_name(value) -> str:
    """
    Normalize an Excel header for easier matching.
    Example:
        'Patient ID' -> 'patientid'
        'Eye_ID'     -> 'eyeid'
    """
    return re.sub(r"[^a-z0-9]", "", safe_text(value).lower())


def natural_key(text: str):
    """
    Natural sorting:
        1_2.jpg before 1_10.jpg
    """
    return [
        int(part) if part.isdigit() else part.lower()
        for part in re.split(r"(\d+)", text)
    ]


def extract_prefix_id(filename: str) -> str | None:
    """
    For filenames such as:
        72_10.jpg
        1_2.json

    returns:
        '72'
        '1'

    IMPORTANT:
    We initially call this a 'prefix ID'.
    We do NOT assume it is an Eye ID until verified against demographics.
    """
    stem = Path(filename).stem

    match = re.match(r"^([^_]+)_", stem)

    if match:
        return match.group(1)

    return None


def read_split_file(path: Path) -> list[str]:
    if not path.exists():
        return []

    names = []

    with path.open("r", encoding="utf-8-sig") as f:
        for line in f:
            line = line.strip()

            if line:
                names.append(Path(line).name)

    return names


def find_matching_file_by_stem(
    directory: Path,
    stem: str,
    allowed_extensions: set[str]
) -> Path | None:

    if not directory.exists():
        return None

    for p in directory.iterdir():
        if (
            p.is_file()
            and p.stem == stem
            and p.suffix.lower() in allowed_extensions
        ):
            return p

    return None


def write_csv(path: Path, rows: list[dict], fieldnames=None):
    path.parent.mkdir(parents=True, exist_ok=True)

    if not rows:
        path.write_text("", encoding="utf-8")
        return

    if fieldnames is None:
        fieldnames = list(rows[0].keys())

    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def json_dump(path: Path, data):
    path.parent.mkdir(parents=True, exist_ok=True)

    with path.open("w", encoding="utf-8") as f:
        json.dump(
            data,
            f,
            indent=2,
            ensure_ascii=False,
            default=str
        )


# ============================================================
# DEMOGRAPHICS AUDIT
# ============================================================

def find_demographics_file():
    candidates = sorted(
        DATASET_ROOT.glob("*.xlsx"),
        key=lambda p: p.name.lower()
    )

    for p in candidates:
        if "demograph" in p.name.lower():
            return p

    return candidates[0] if candidates else None


def audit_demographics():
    result = {
        "file": None,
        "available": False,
        "headers": [],
        "rows": [],
        "eye_to_patient": {},
        "error": None,
    }

    xlsx_path = find_demographics_file()

    if xlsx_path is None:
        result["error"] = "No XLSX demographics file found."
        return result

    result["file"] = str(xlsx_path)

    if load_workbook is None:
        result["error"] = (
            "openpyxl is not installed. "
            "Run: pip install openpyxl"
        )
        return result

    try:
        wb = load_workbook(
            xlsx_path,
            read_only=True,
            data_only=True
        )

        ws = wb.active

        rows = list(ws.iter_rows(values_only=True))

        if not rows:
            result["error"] = "Excel file is empty."
            return result

        headers = [safe_text(v) for v in rows[0]]
        normalized_headers = [
            normalize_column_name(v)
            for v in headers
        ]

        result["headers"] = headers

        patient_col = None
        eye_col = None
        age_col = None
        gender_col = None
        eye_category_col = None

        for idx, col in enumerate(normalized_headers):

            if col in {"patientid", "patient"}:
                patient_col = idx

            elif col in {"eyeid", "eye"}:
                eye_col = idx

            elif col.startswith("age"):
                age_col = idx

            elif "gender" in col or col == "sex":
                gender_col = idx

            elif (
                "eyecateg" in col
                or "eyecategory" in col
                or col in {"laterality", "eyeside"}
            ):
                eye_category_col = idx

        parsed_rows = []
        eye_to_patient = {}

        for row_number, row in enumerate(rows[1:], start=2):

            values = list(row)

            if not any(v is not None for v in values):
                continue

            record = {
                "excel_row": row_number,
                "patient_id": (
                    safe_text(values[patient_col])
                    if patient_col is not None
                    and patient_col < len(values)
                    else ""
                ),
                "eye_id": (
                    safe_text(values[eye_col])
                    if eye_col is not None
                    and eye_col < len(values)
                    else ""
                ),
                "eye_category": (
                    safe_text(values[eye_category_col])
                    if eye_category_col is not None
                    and eye_category_col < len(values)
                    else ""
                ),
                "age": (
                    values[age_col]
                    if age_col is not None
                    and age_col < len(values)
                    else None
                ),
                "gender": (
                    safe_text(values[gender_col])
                    if gender_col is not None
                    and gender_col < len(values)
                    else ""
                ),
            }

            parsed_rows.append(record)

            if record["eye_id"]:
                eye_to_patient[record["eye_id"]] = record["patient_id"]

        result["available"] = True
        result["rows"] = parsed_rows
        result["eye_to_patient"] = eye_to_patient

        wb.close()

    except Exception as e:
        result["error"] = repr(e)

    return result


# ============================================================
# RAW LABELME AUDIT
# ============================================================

def audit_labelme():
    print("\n[1/5] Auditing raw LabelMe annotations...")

    result = {
        "json_count": 0,
        "image_count": 0,
        "folders": set(),
        "label_counts": Counter(),
        "label_image_counts": Counter(),
        "shape_type_counts": Counter(),
        "label_shape_counts": Counter(),
        "label_point_counts": defaultdict(list),
        "combination_counts": Counter(),
        "json_errors": [],
        "pairing_issues": [],
        "annotation_details": [],
        "image_level": [],
        "image_dimensions": Counter(),
        "json_versions": Counter(),
    }

    if not RAW_LABELME_DIR.exists():
        raise FileNotFoundError(
            f"Raw LabelMe folder not found:\n{RAW_LABELME_DIR}"
        )

    json_files = sorted(
        RAW_LABELME_DIR.rglob("*.json"),
        key=lambda p: natural_key(str(p))
    )

    image_files = sorted(
        [
            p
            for p in RAW_LABELME_DIR.rglob("*")
            if p.is_file()
            and p.suffix.lower() in IMAGE_EXTENSIONS
        ],
        key=lambda p: natural_key(str(p))
    )

    result["json_count"] = len(json_files)
    result["image_count"] = len(image_files)

    image_lookup = {
        str(p.relative_to(RAW_LABELME_DIR).with_suffix("")): p
        for p in image_files
    }

    json_lookup = {
        str(p.relative_to(RAW_LABELME_DIR).with_suffix("")): p
        for p in json_files
    }

    # --------------------------------------------------------
    # Pairing checks
    # --------------------------------------------------------

    all_pair_keys = set(image_lookup) | set(json_lookup)

    for key in sorted(all_pair_keys, key=natural_key):

        image_exists = key in image_lookup
        json_exists = key in json_lookup

        if image_exists != json_exists:
            result["pairing_issues"].append({
                "relative_stem": key,
                "image_exists": image_exists,
                "json_exists": json_exists,
                "image_path": (
                    str(image_lookup[key])
                    if image_exists
                    else ""
                ),
                "json_path": (
                    str(json_lookup[key])
                    if json_exists
                    else ""
                ),
            })

    # --------------------------------------------------------
    # Parse JSONs
    # --------------------------------------------------------

    for json_path in json_files:

        relative_json = json_path.relative_to(RAW_LABELME_DIR)
        relative_stem = relative_json.with_suffix("")

        result["folders"].add(str(relative_json.parent))

        try:
            with json_path.open(
                "r",
                encoding="utf-8-sig"
            ) as f:
                data = json.load(f)

        except Exception as e:
            result["json_errors"].append({
                "json_file": str(relative_json),
                "error": repr(e),
            })
            continue

        version = safe_text(data.get("version"))

        if version:
            result["json_versions"][version] += 1

        width = data.get("imageWidth")
        height = data.get("imageHeight")

        if width is not None and height is not None:
            result["image_dimensions"][
                f"{width}x{height}"
            ] += 1

        shapes = data.get("shapes", [])

        if not isinstance(shapes, list):
            result["json_errors"].append({
                "json_file": str(relative_json),
                "error": "'shapes' is not a list",
            })
            continue

        image_labels = set()
        shape_count_for_image = 0

        for shape_index, shape in enumerate(shapes):

            if not isinstance(shape, dict):
                result["json_errors"].append({
                    "json_file": str(relative_json),
                    "error": (
                        f"Shape #{shape_index} "
                        f"is not a dictionary"
                    ),
                })
                continue

            label = safe_text(
                shape.get("label")
            ) or "__MISSING_LABEL__"

            shape_type = safe_text(
                shape.get("shape_type")
            ) or "__MISSING_SHAPE_TYPE__"

            points = shape.get("points", [])

            if not isinstance(points, list):
                points = []

            n_points = len(points)

            valid_coordinate_pairs = 0
            invalid_points = 0

            for point in points:

                if (
                    isinstance(point, (list, tuple))
                    and len(point) >= 2
                    and isinstance(point[0], (int, float))
                    and isinstance(point[1], (int, float))
                ):
                    valid_coordinate_pairs += 1
                else:
                    invalid_points += 1

            result["label_counts"][label] += 1
            result["shape_type_counts"][shape_type] += 1
            result["label_shape_counts"][
                (label, shape_type)
            ] += 1

            result["label_point_counts"][label].append(
                n_points
            )

            image_labels.add(label)
            shape_count_for_image += 1

            result["annotation_details"].append({
                "folder": str(relative_json.parent),
                "json_file": relative_json.name,
                "image_stem": relative_stem.name,
                "prefix_id": extract_prefix_id(
                    relative_stem.name
                ),
                "shape_index": shape_index,
                "label": label,
                "shape_type": shape_type,
                "n_points": n_points,
                "valid_coordinate_pairs":
                    valid_coordinate_pairs,
                "invalid_points": invalid_points,
                "group_id": safe_text(
                    shape.get("group_id")
                ),
                "flags": json.dumps(
                    shape.get("flags", {}),
                    ensure_ascii=False
                ),
            })

        for label in image_labels:
            result["label_image_counts"][label] += 1

        sorted_labels = tuple(
            sorted(image_labels)
        )

        combination_name = (
            " + ".join(sorted_labels)
            if sorted_labels
            else "__NO_ANNOTATIONS__"
        )

        result["combination_counts"][
            combination_name
        ] += 1

        result["image_level"].append({
            "folder": str(relative_json.parent),
            "json_file": relative_json.name,
            "image_stem": relative_stem.name,
            "prefix_id": extract_prefix_id(
                relative_stem.name
            ),
            "image_width": width,
            "image_height": height,
            "n_annotations": shape_count_for_image,
            "n_unique_labels": len(image_labels),
            "labels": ";".join(sorted(image_labels)),
        })

    return result


# ============================================================
# OFFICIAL SEGMENTATION SPLIT AUDIT
# ============================================================

def audit_official_split(demographics):
    print("[2/5] Auditing Image_Seg and official train/test split...")

    result = {
        "images_count": 0,
        "masks_count": 0,
        "train_count": 0,
        "test_count": 0,
        "train_duplicates": [],
        "test_duplicates": [],
        "filename_overlap": [],
        "prefix_overlap": [],
        "patient_overlap": [],
        "missing_split_images": [],
        "missing_split_masks": [],
        "split_rows": [],
        "prefix_interpreted_as_eye_id": False,
    }

    images_dir = IMAGE_SEG_DIR / "images"
    masks_dir = IMAGE_SEG_DIR / "masks"

    train_file = IMAGE_SEG_DIR / "train.txt"
    test_file = IMAGE_SEG_DIR / "test.txt"

    if images_dir.exists():
        images = [
            p
            for p in images_dir.iterdir()
            if p.is_file()
            and p.suffix.lower() in IMAGE_EXTENSIONS
        ]
    else:
        images = []

    if masks_dir.exists():
        masks = [
            p
            for p in masks_dir.iterdir()
            if p.is_file()
            and p.suffix.lower() in IMAGE_EXTENSIONS
        ]
    else:
        masks = []

    result["images_count"] = len(images)
    result["masks_count"] = len(masks)

    train_names = read_split_file(train_file)
    test_names = read_split_file(test_file)

    result["train_count"] = len(train_names)
    result["test_count"] = len(test_names)

    train_counter = Counter(train_names)
    test_counter = Counter(test_names)

    result["train_duplicates"] = sorted(
        [
            name
            for name, count in train_counter.items()
            if count > 1
        ],
        key=natural_key
    )

    result["test_duplicates"] = sorted(
        [
            name
            for name, count in test_counter.items()
            if count > 1
        ],
        key=natural_key
    )

    train_set = set(train_names)
    test_set = set(test_names)

    result["filename_overlap"] = sorted(
        train_set & test_set,
        key=natural_key
    )

    train_prefixes = {
        extract_prefix_id(name)
        for name in train_names
        if extract_prefix_id(name) is not None
    }

    test_prefixes = {
        extract_prefix_id(name)
        for name in test_names
        if extract_prefix_id(name) is not None
    }

    result["prefix_overlap"] = sorted(
        train_prefixes & test_prefixes,
        key=natural_key
    )

    # --------------------------------------------------------
    # Check whether filename prefix corresponds to Eye ID
    # --------------------------------------------------------

    eye_to_patient = demographics.get(
        "eye_to_patient",
        {}
    )

    all_prefixes = train_prefixes | test_prefixes

    if eye_to_patient and all_prefixes:

        matching_prefixes = sum(
            1
            for prefix in all_prefixes
            if prefix in eye_to_patient
        )

        # Only interpret prefix as Eye ID if every split prefix
        # maps to an Eye ID in the demographics table.
        result["prefix_interpreted_as_eye_id"] = (
            matching_prefixes == len(all_prefixes)
        )

    train_patients = set()
    test_patients = set()

    for split_name, filenames in [
        ("train", train_names),
        ("test", test_names),
    ]:

        for filename in filenames:

            stem = Path(filename).stem
            prefix_id = extract_prefix_id(filename)

            image_path = find_matching_file_by_stem(
                images_dir,
                stem,
                IMAGE_EXTENSIONS
            )

            mask_path = find_matching_file_by_stem(
                masks_dir,
                stem,
                IMAGE_EXTENSIONS
            )

            if image_path is None:
                result["missing_split_images"].append({
                    "split": split_name,
                    "filename": filename,
                })

            if mask_path is None:
                result["missing_split_masks"].append({
                    "split": split_name,
                    "filename": filename,
                })

            eye_id = ""
            patient_id = ""

            if (
                result["prefix_interpreted_as_eye_id"]
                and prefix_id is not None
            ):
                eye_id = prefix_id
                patient_id = eye_to_patient.get(
                    eye_id,
                    ""
                )

                if patient_id:
                    if split_name == "train":
                        train_patients.add(patient_id)
                    else:
                        test_patients.add(patient_id)

            result["split_rows"].append({
                "split": split_name,
                "filename": filename,
                "prefix_id": prefix_id or "",
                "interpreted_eye_id": eye_id,
                "patient_id": patient_id,
                "image_exists": image_path is not None,
                "mask_exists": mask_path is not None,
            })

    result["patient_overlap"] = sorted(
        train_patients & test_patients,
        key=natural_key
    )

    result["train_unique_prefixes"] = len(
        train_prefixes
    )

    result["test_unique_prefixes"] = len(
        test_prefixes
    )

    result["train_unique_patients"] = len(
        train_patients
    )

    result["test_unique_patients"] = len(
        test_patients
    )

    return result


# ============================================================
# MASK VALUE AUDIT
# ============================================================

def audit_masks():
    print("[3/5] Auditing segmentation masks...")

    result = {
        "available": False,
        "mask_count": 0,
        "mode_counts": Counter(),
        "dimension_counts": Counter(),
        "global_pixel_values": Counter(),
        "per_mask_rows": [],
        "errors": [],
        "warning": None,
    }

    masks_dir = IMAGE_SEG_DIR / "masks"

    if not masks_dir.exists():
        result["warning"] = "Masks directory not found."
        return result

    mask_files = sorted(
        [
            p
            for p in masks_dir.iterdir()
            if p.is_file()
            and p.suffix.lower() in IMAGE_EXTENSIONS
        ],
        key=lambda p: natural_key(p.name)
    )

    result["mask_count"] = len(mask_files)

    if Image is None:
        result["warning"] = (
            "Pillow is not installed. "
            "Mask pixel-value audit skipped. "
            "Run: pip install pillow"
        )
        return result

    result["available"] = True

    for mask_path in mask_files:

        try:
            with Image.open(mask_path) as img:

                result["mode_counts"][img.mode] += 1
                result["dimension_counts"][
                    f"{img.width}x{img.height}"
                ] += 1

                values = Counter(img.getdata())

                for pixel_value, count in values.items():
                    result["global_pixel_values"][
                        str(pixel_value)
                    ] += count

                result["per_mask_rows"].append({
                    "mask_file": mask_path.name,
                    "mode": img.mode,
                    "width": img.width,
                    "height": img.height,
                    "n_unique_values": len(values),
                    "unique_values": ";".join(
                        sorted(
                            [str(v) for v in values.keys()]
                        )
                    ),
                })

        except Exception as e:
            result["errors"].append({
                "mask_file": mask_path.name,
                "error": repr(e),
            })

    return result


# ============================================================
# HF FOLDER AUDIT
# ============================================================

def audit_hf_folder():
    print("[4/5] Auditing RVO_HF_Detection folder...")

    result = {
        "exists": HF_DETECTION_DIR.exists(),
        "total_files": 0,
        "image_files": 0,
        "extension_counts": Counter(),
    }

    if not HF_DETECTION_DIR.exists():
        return result

    files = [
        p
        for p in HF_DETECTION_DIR.rglob("*")
        if p.is_file()
    ]

    result["total_files"] = len(files)

    for p in files:

        ext = p.suffix.lower() or "__NO_EXTENSION__"

        result["extension_counts"][ext] += 1

        if ext in IMAGE_EXTENSIONS:
            result["image_files"] += 1

    return result


# ============================================================
# DEMOGRAPHIC SUMMARY
# ============================================================

def demographic_summary(demographics):
    rows = demographics.get("rows", [])

    result = {
        "rows": len(rows),
        "unique_patients": 0,
        "unique_eyes": 0,
        "gender_counts": {},
        "eye_category_counts": {},
        "age_summary": {},
    }

    if not rows:
        return result

    patients = {
        r["patient_id"]
        for r in rows
        if r["patient_id"]
    }

    eyes = {
        r["eye_id"]
        for r in rows
        if r["eye_id"]
    }

    genders = Counter(
        r["gender"]
        for r in rows
        if r["gender"]
    )

    eye_categories = Counter(
        r["eye_category"]
        for r in rows
        if r["eye_category"]
    )

    ages = []

    for r in rows:
        try:
            if r["age"] is not None:
                ages.append(float(r["age"]))
        except (TypeError, ValueError):
            pass

    result["unique_patients"] = len(patients)
    result["unique_eyes"] = len(eyes)
    result["gender_counts"] = dict(genders)
    result["eye_category_counts"] = dict(
        eye_categories
    )

    if ages:
        result["age_summary"] = {
            "count": len(ages),
            "minimum": min(ages),
            "maximum": max(ages),
            "mean": mean(ages),
            "median": median(ages),
        }

    return result


# ============================================================
# SAVE REPORTS
# ============================================================

def save_reports(
    demographics,
    raw,
    split,
    masks,
    hf
):
    print("[5/5] Saving audit reports...")

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    # --------------------------------------------------------
    # Label × shape type
    # --------------------------------------------------------

    label_shape_rows = []

    for (label, shape_type), count in sorted(
        raw["label_shape_counts"].items()
    ):
        label_shape_rows.append({
            "label": label,
            "shape_type": shape_type,
            "annotation_count": count,
        })

    write_csv(
        OUTPUT_DIR / "label_shape_counts.csv",
        label_shape_rows
    )

    # --------------------------------------------------------
    # Label summary
    # --------------------------------------------------------

    label_summary_rows = []

    for label in sorted(raw["label_counts"]):

        point_counts = raw[
            "label_point_counts"
        ].get(label, [])

        label_summary_rows.append({
            "label": label,
            "annotation_count":
                raw["label_counts"][label],
            "images_containing_label":
                raw["label_image_counts"][label],
            "min_points_per_annotation":
                min(point_counts)
                if point_counts else "",
            "max_points_per_annotation":
                max(point_counts)
                if point_counts else "",
            "mean_points_per_annotation":
                (
                    round(mean(point_counts), 3)
                    if point_counts
                    else ""
                ),
            "median_points_per_annotation":
                (
                    median(point_counts)
                    if point_counts
                    else ""
                ),
        })

    write_csv(
        OUTPUT_DIR / "label_summary.csv",
        label_summary_rows
    )

    # --------------------------------------------------------
    # Shape-type summary
    # --------------------------------------------------------

    write_csv(
        OUTPUT_DIR / "shape_type_counts.csv",
        [
            {
                "shape_type": shape_type,
                "annotation_count": count,
            }
            for shape_type, count
            in sorted(
                raw["shape_type_counts"].items()
            )
        ]
    )

    # --------------------------------------------------------
    # Annotation details
    # --------------------------------------------------------

    write_csv(
        OUTPUT_DIR / "annotation_details.csv",
        raw["annotation_details"]
    )

    # --------------------------------------------------------
    # Image-level annotations
    # --------------------------------------------------------

    write_csv(
        OUTPUT_DIR / "image_level_annotations.csv",
        raw["image_level"]
    )

    # --------------------------------------------------------
    # Annotation combinations
    # --------------------------------------------------------

    write_csv(
        OUTPUT_DIR / "biomarker_combinations.csv",
        [
            {
                "combination": combination,
                "image_count": count,
            }
            for combination, count
            in raw["combination_counts"].most_common()
        ]
    )

    # --------------------------------------------------------
    # Pairing issues
    # --------------------------------------------------------

    write_csv(
        OUTPUT_DIR / "pairing_issues.csv",
        raw["pairing_issues"]
    )

    write_csv(
        OUTPUT_DIR / "json_errors.csv",
        raw["json_errors"]
    )

    # --------------------------------------------------------
    # Official split
    # --------------------------------------------------------

    write_csv(
        OUTPUT_DIR / "official_split_files.csv",
        split["split_rows"]
    )

    write_csv(
        OUTPUT_DIR / "missing_split_images.csv",
        split["missing_split_images"]
    )

    write_csv(
        OUTPUT_DIR / "missing_split_masks.csv",
        split["missing_split_masks"]
    )

    # --------------------------------------------------------
    # Demographics
    # --------------------------------------------------------

    write_csv(
        OUTPUT_DIR / "demographics_parsed.csv",
        demographics.get("rows", [])
    )

    # --------------------------------------------------------
    # Mask details
    # --------------------------------------------------------

    write_csv(
        OUTPUT_DIR / "mask_details.csv",
        masks["per_mask_rows"]
    )

    write_csv(
        OUTPUT_DIR / "mask_errors.csv",
        masks["errors"]
    )

    # --------------------------------------------------------
    # Comprehensive JSON summary
    # --------------------------------------------------------

    summary = {
        "dataset_root": str(DATASET_ROOT),

        "raw_labelme": {
            "json_count": raw["json_count"],
            "image_count": raw["image_count"],
            "folder_count": len(raw["folders"]),
            "folders": sorted(raw["folders"]),
            "label_counts":
                dict(raw["label_counts"]),
            "label_image_counts":
                dict(raw["label_image_counts"]),
            "shape_type_counts":
                dict(raw["shape_type_counts"]),
            "label_shape_counts": {
                f"{label} | {shape}": count
                for (label, shape), count
                in raw["label_shape_counts"].items()
            },
            "biomarker_combinations":
                dict(raw["combination_counts"]),
            "image_dimensions":
                dict(raw["image_dimensions"]),
            "labelme_versions":
                dict(raw["json_versions"]),
            "pairing_issue_count":
                len(raw["pairing_issues"]),
            "json_error_count":
                len(raw["json_errors"]),
        },

        "official_split": {
            key: value
            for key, value in split.items()
            if key != "split_rows"
        },

        "demographics": {
            "file":
                demographics.get("file"),
            "available":
                demographics.get("available"),
            "error":
                demographics.get("error"),
            "summary":
                demographic_summary(
                    demographics
                ),
        },

        "masks": {
            "available": masks["available"],
            "mask_count": masks["mask_count"],
            "mode_counts":
                dict(masks["mode_counts"]),
            "dimension_counts":
                dict(masks["dimension_counts"]),
            "global_pixel_values":
                dict(masks["global_pixel_values"]),
            "error_count":
                len(masks["errors"]),
            "warning":
                masks["warning"],
        },

        "hf_detection": {
            "exists": hf["exists"],
            "total_files": hf["total_files"],
            "image_files": hf["image_files"],
            "extension_counts":
                dict(hf["extension_counts"]),
        },
    }

    json_dump(
        OUTPUT_DIR / "audit_summary.json",
        summary
    )

    return summary


# ============================================================
# TERMINAL SUMMARY
# ============================================================

def print_summary(
    demographics,
    raw,
    split,
    masks,
    hf
):
    print("\n")
    print("=" * 72)
    print("RVO-ME DATASET AUDIT SUMMARY")
    print("=" * 72)

    print("\nRAW LABELME DATA")
    print("-" * 72)

    print(
        f"Images:                 "
        f"{raw['image_count']}"
    )

    print(
        f"JSON files:             "
        f"{raw['json_count']}"
    )

    print(
        f"Folders:                "
        f"{len(raw['folders'])}"
    )

    print(
        f"Pairing issues:         "
        f"{len(raw['pairing_issues'])}"
    )

    print(
        f"JSON errors:            "
        f"{len(raw['json_errors'])}"
    )

    print("\nLABEL COUNTS")
    print("-" * 72)

    for label, count in sorted(
        raw["label_counts"].items()
    ):
        image_count = raw[
            "label_image_counts"
        ][label]

        print(
            f"{label:<20} "
            f"annotations={count:<8} "
            f"images={image_count}"
        )

    print("\nLABEL × SHAPE TYPE")
    print("-" * 72)

    for (label, shape), count in sorted(
        raw["label_shape_counts"].items()
    ):
        print(
            f"{label:<20} "
            f"{shape:<20} "
            f"{count}"
        )

    print("\nIMAGE DIMENSIONS")
    print("-" * 72)

    for dim, count in raw[
        "image_dimensions"
    ].most_common():
        print(
            f"{dim:<20} {count}"
        )

    print("\nOFFICIAL IMAGE_SEG SPLIT")
    print("-" * 72)

    print(
        f"Images directory:       "
        f"{split['images_count']}"
    )

    print(
        f"Masks directory:        "
        f"{split['masks_count']}"
    )

    print(
        f"Train entries:          "
        f"{split['train_count']}"
    )

    print(
        f"Test entries:           "
        f"{split['test_count']}"
    )

    print(
        f"Exact filename overlap: "
        f"{len(split['filename_overlap'])}"
    )

    print(
        f"Prefix ID overlap:      "
        f"{len(split['prefix_overlap'])}"
    )

    print(
        "Prefix maps to Eye ID:  "
        f"{split['prefix_interpreted_as_eye_id']}"
    )

    if split[
        "prefix_interpreted_as_eye_id"
    ]:

        print(
            f"Train patients:         "
            f"{split['train_unique_patients']}"
        )

        print(
            f"Test patients:          "
            f"{split['test_unique_patients']}"
        )

        print(
            f"Patient overlap:        "
            f"{len(split['patient_overlap'])}"
        )

        if split["patient_overlap"]:
            print(
                "WARNING: patient IDs occur "
                "in both train and test!"
            )

    print("\nDEMOGRAPHICS")
    print("-" * 72)

    demo_summary = demographic_summary(
        demographics
    )

    print(
        f"Available:              "
        f"{demographics.get('available')}"
    )

    print(
        f"Unique patients:        "
        f"{demo_summary['unique_patients']}"
    )

    print(
        f"Unique eyes:            "
        f"{demo_summary['unique_eyes']}"
    )

    print(
        f"Gender counts:          "
        f"{demo_summary['gender_counts']}"
    )

    print(
        f"Eye categories:         "
        f"{demo_summary['eye_category_counts']}"
    )

    print("\nMASK AUDIT")
    print("-" * 72)

    print(
        f"Masks:                  "
        f"{masks['mask_count']}"
    )

    print(
        f"Pixel audit available:  "
        f"{masks['available']}"
    )

    if masks["available"]:
        print(
            f"Mask modes:             "
            f"{dict(masks['mode_counts'])}"
        )

        print(
            f"Global pixel values:    "
            f"{dict(masks['global_pixel_values'])}"
        )

    elif masks["warning"]:
        print(
            f"Note:                   "
            f"{masks['warning']}"
        )

    print("\nHF DETECTION")
    print("-" * 72)

    print(
        f"Folder exists:          "
        f"{hf['exists']}"
    )

    print(
        f"Files:                  "
        f"{hf['total_files']}"
    )

    print(
        f"Images:                 "
        f"{hf['image_files']}"
    )

    print("\n")
    print("=" * 72)
    print("AUDIT COMPLETE")
    print("=" * 72)

    print(
        "\nReports saved to:\n"
        f"{OUTPUT_DIR}"
    )

    print(
        "\nMost important files to send/review next:"
    )

    print(
        "  1. audit_summary.json"
    )

    print(
        "  2. label_shape_counts.csv"
    )

    print(
        "  3. label_summary.csv"
    )

    print(
        "  4. biomarker_combinations.csv"
    )

    print(
        "  5. pairing_issues.csv"
    )

    print(
        "  6. official_split_files.csv"
    )


# ============================================================
# MAIN
# ============================================================

def main():

    print("=" * 72)
    print("RVO-ME DATASET AUDIT")
    print("=" * 72)

    print(
        f"\nDataset root:\n"
        f"{DATASET_ROOT}"
    )

    if not DATASET_ROOT.exists():
        print(
            "\nERROR: Dataset root does not exist."
        )
        sys.exit(1)

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    demographics = audit_demographics()

    raw = audit_labelme()

    split = audit_official_split(
        demographics
    )

    masks = audit_masks()

    hf = audit_hf_folder()

    save_reports(
        demographics,
        raw,
        split,
        masks,
        hf
    )

    print_summary(
        demographics,
        raw,
        split,
        masks,
        hf
    )


if __name__ == "__main__":
    main()